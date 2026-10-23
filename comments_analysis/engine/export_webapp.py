"""Write an analysis run as import files for the Workload Forecast web application.

    planning.xlsx    Planning Input (one row per department / sub-department / category / project)
    proposals.xlsx   department feedback: only the months the department changed; blank = keep the planner value;
                     rows without a planning line become new-allocation requests in the application. A second
                     sheet "Reviewed departments" lists every department whose feedback file was processed, so
                     departments that answered without a man-hour change are recorded as "no change".
    after.xlsx       after-implementation extract
    departments.csv  department master with headcount (capacity)

Load them in the application under Cycle Control > Import data (and Administration > Reference data), or
let run_analysis.py --webapp-export write them for every cycle.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from analysis import DEPARTMENTS
from compare import KEY

COLUMNS = ["Dept", "SubDept", "Category", "ProjectCode", "Activity", "Planner", "Reason", "Occurrence", "RequestedCategory"]
PARENT = {"2410": "2400", "2420": "2400", "2430": "2400", "2440": "2400", "2450": "2400", "2710": "2700", "2720": "2700"}


def _write(path: Path, rows: list[list], months: list[dt.date], reviewed: list[str] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(COLUMNS + [dt.datetime.combine(m, dt.time()) for m in months])
    for c in ws[1]:
        c.font = Font(bold=True)
        if isinstance(c.value, dt.datetime):
            c.number_format = "mmm-yy"
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "J2"
    if reviewed is not None:
        ls = wb.create_sheet("Reviewed departments")
        ls.append(["Dept"])
        ls["A1"].font = Font(bold=True)
        for d in reviewed:
            ls.append([d])
    wb.save(path)


def export_webapp(folder: Path, *, rows: pd.DataFrame, cells: pd.DataFrame, cap: pd.DataFrame,
                  fb: pd.DataFrame | None = None) -> dict[str, int]:
    folder.mkdir(parents=True, exist_ok=True)
    months = sorted(cells["Month"].unique())
    piv = {col: cells.pivot_table(index=KEY, columns="Month", values=col, aggfunc="sum").reindex(columns=months)
           for col in ("P0", "Final", "I0")}
    changed = cells[cells["Changed"]].set_index(KEY + ["Month"])["Final"]
    r = rows.copy()
    # the row's own project code (a department's code correction is paired to the planner's row, whose code stays)
    r["Code"] = r["ProjectFull"].where(r["ProjectFull"].fillna("") != "", r["ProjectKey"].str.split("#").str[0])
    r["Occ"] = r["ProjectKey"].str.extract(r"#(\d+)$")[0].fillna("1").astype(int)
    r = r.sort_values(["DeptCode", "SectionCode", "Category", "Code", "Occ"])

    def base(x, name):
        # Occurrence keeps two rows with the same code apart even when a file lists only some rows
        return [x.DeptCode, x.SectionCode, x.Category, x.Code, str(name or "")[:250], "", "", int(x.Occ), ""]

    planning, proposals, after = [], [], []
    for _, x in r.iterrows():
        key = tuple(x[k] for k in KEY)
        if x.P_Exists:
            planning.append(base(x, x.P_Name) + [float(v) if v else None for v in piv["P0"].loc[key]])
        if x.I_Exists:
            after.append(base(x, x.I_Name) + [float(v) for v in piv["I0"].loc[key]])
        if x.D_Exists and x.Concern:
            reason = x.D_Remark if isinstance(x.D_Remark, str) and x.D_Remark.strip() else (
                f"Imported from {x.D_File} row {int(x.D_Row)}" if pd.notna(x.D_Row) else "Imported from the department feedback file")
            row = base(x, x.D_Name)
            if isinstance(x.CodeCorrection, str) and x.CodeCorrection:
                # the application has no code-change request: the requested code travels in the reason
                old, new = x.CodeCorrection.split(" (")[0].split(" -> ")
                reason = (f"{reason} | Department {'typed' if 'typed' in x.CodeCorrection else 'used'} project code {new} "
                          f"for {old} - confirm the code in Primavera")
            row[6] = reason[:2000]
            if x.P_Exists:
                vals = [float(changed.loc[key + (m,)]) if key + (m,) in changed.index else None for m in months]
            else:
                vals = [float(v) if v else None for v in piv["Final"].loc[key]]
            # a project the department filed under another phase becomes a category-change request
            row[8] = x.RequestedPhase if isinstance(x.RequestedPhase, str) else ""
            if any(v is not None for v in vals) or row[8]:
                proposals.append(row + vals)
    _write(folder / "planning.xlsx", planning, months)
    reviewed = sorted(fb.loc[fb["Feedback_Processed"].astype(bool), "DeptCode"].astype(str), key=lambda c: (len(c), c)) \
        if fb is not None and "Feedback_Processed" in fb else None
    _write(folder / "proposals.xlsx", proposals, months, reviewed)
    _write(folder / "after.xlsx", after, months)
    d = cap.rename(columns={"DeptCode": "Code", "LoanTo_Unpaid": "LoanToUnpaid"}).copy()
    d["Name"] = d["Code"].map(DEPARTMENTS)
    d["ParentCode"] = d["Code"].map(PARENT).fillna("")
    d[["Code", "Name", "ParentCode", "Permanent", "Temporary", "Managers", "LoanToUnpaid"]].fillna(0).to_csv(
        folder / "departments.csv", index=False)
    return {"planning": len(planning), "proposals": len(proposals), "after": len(after), "first_month": str(months[0]),
            "months": len(months), "reviewed_departments": len(reviewed or [])}
