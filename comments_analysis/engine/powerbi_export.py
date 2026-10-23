"""Star-schema tables for the 'Workload Comments Analysis' Power BI report.

Writes CSV files plus one Excel workbook (one formatted Excel Table per sheet) so the model can be
loaded with Get Data > Excel; column names match across tables, so Power BI auto-detects the joins.
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from analysis import CATEGORY_ORDER, DEPARTMENTS, FIRM_CATEGORIES, HOURS_PER_MAN_MONTH
from compare import BRIEF_MAP, KEY

PACKAGE = Path(__file__).resolve().parent.parent / "powerbi"     # queries, measures, theme and guide shipped with the data

SCENARIOS = [("Planning Input", 1, "#2A78D6", "Man-hours distributed by Planning and sent to departments"),
             ("Discipline Comment", 2, "#EB6834", "Planning Input with populated, different department values applied (blank = retained)"),
             ("After Implementation", 3, "#1BAF7A", "Re-extracted forecast after planners implemented comments")]
CATEGORY_COLORS = {"Running": "#2E7D32", "Internal / Unforeseen": "#00897B", "Loan To": "#607D8B",
                   "Under Signature": "#6A1B9A", "Proposal": "#EF6C00", "Potential Opportunity": "#F9A825"}
DIVISION = {"2410": "2400", "2420": "2400", "2430": "2400", "2440": "2400", "2450": "2400", "2710": "2700", "2720": "2700"}
SERVICE = {"300": "Management", "600": "Construction", "2110": "Management", "3100": "Management", "3300": "Management",
           "3500": "Management", "3400": "Construction", "4100": "Construction", "5000": "Procurement", "5100": "Procurement"}


def _row_key(df):
    return df["DeptCode"].astype(str) + "|" + df["SectionCode"].astype(str) + "|" + df["Phase"] + "|" + df["ProjectKey"]


def _client(name: str) -> str:
    s = re.sub(r"^\s*[\d\-\s,_.]+", "", str(name or ""))
    return s.split(",")[0].strip()[:40] if "," in s else ""


def export_powerbi(folder: Path, *, df, cells, rows, ps, cap, issues, cycle: str, possible=None, fb=None, inventory=None):
    folder.mkdir(parents=True, exist_ok=True)
    c = cells.copy()
    c["RowKey"] = _row_key(c)
    r = rows.copy()
    r["RowKey"] = _row_key(r)
    cat = r.set_index("RowKey")["Category"]
    proj = r.set_index("RowKey")["ProjectFull"]
    long = []
    for scen, col in (("Planning Input", "P0"), ("Discipline Comment", "Final"), ("After Implementation", "I0")):
        x = c[c[col].abs() > 0][["RowKey", "DeptCode", "Month", col]].rename(columns={col: "MH"})
        x["Scenario"] = scen
        long.append(x)
    fact = pd.concat(long, ignore_index=True)
    fact["Category"] = fact["RowKey"].map(cat)
    fact["ProjectCode"] = fact["RowKey"].map(proj)
    fact["Cycle"] = cycle
    fact["DeptCode"] = fact["DeptCode"].astype(int)
    fact["Month"] = pd.to_datetime(fact["Month"])
    fact = fact[["Cycle", "Scenario", "DeptCode", "Category", "ProjectCode", "RowKey", "Month", "MH"]]

    sev = ps.set_index("ProjectFull")
    items = pd.DataFrame({
        "Cycle": cycle, "RowKey": r["RowKey"], "DeptCode": r["DeptCode"].astype(int), "SubDept": r["SectionCode"],
        "Category": r["Category"], "ProjectCode": r["ProjectFull"],
        "Activity": r["D_Name"].fillna(r["P_Name"]).fillna(r["I_Name"]),
        "Item Planning MH": r["P_Total"], "Item Discipline MH": r["Final_Total"], "Item Implemented MH": r["I_Total"],
        "Gross Increase": r["Gross_Increase"], "Gross Reduction": r["Gross_Reduction"], "Net Change": r["Net_Change"],
        "Concern": r["Concern"].map({True: "Yes", False: "No"}), "Primary Type": r["Primary_Type"],
        "Concern Types": r["Concern_Types"], "Implementation Status": r["Implementation_Status"],
        "Item Implementation Rate": r["Impl_Ratio"], "Implemented Request MH": r["Implemented_Request_MH"],
        "Unrequested Change MH": r["Unrequested_Gross"], "Requested Phase": r["RequestedPhase"],
        "Code Correction": r["CodeCorrection"], "Months Changed": r["Months_Changed"],
        "Blank Not Applied MH": r["Ambiguous_MH"], "Feedback File": r["D_File"], "Feedback Row": r["D_Row"],
        "Brief Categories": r["Brief_Categories"], "Department Remark": r["D_Remark"].fillna(""),
        "Blank Retained MH": r["BlankRetained_MH"], "Bridge New Work MH": r["Bridge_New_Work"],
        "Bridge Increase MH": r["Bridge_Increase"], "Bridge Extension MH": r["Bridge_Extension"],
        "Bridge Reduction MH": r["Bridge_Reduction"], "Bridge Removed MH": r["Bridge_Removed"],
        "Master Flag": r["Master_Flag"] if "Master_Flag" in r else ""})
    # the project's severity in this cycle, on every item, so a report over several cycles rates each cycle on its own
    sev_level = r["ProjectFull"].map(sev["Severity"]).fillna("No Concern")
    items["Project Severity"] = sev_level
    items["Project Severity Points"] = r["ProjectFull"].map(sev["Severity_Points"]).fillna(0)
    items["Project Severity Order"] = sev_level.map({"Critical": 1, "High": 2, "Medium": 3, "Low": 4, "No Concern": 5})
    bridge = (items[["Cycle", "RowKey", "Concern Types"]].assign(ConcernType=items["Concern Types"].str.split("; "))
              .explode("ConcernType").query("ConcernType != ''").dropna()[["Cycle", "RowKey", "ConcernType"]])
    bridge["Brief Category"] = bridge["ConcernType"].map(BRIEF_MAP).fillna("Other")

    ch = c[c["Changed"] | c["AmbiguousBlank"]].copy()
    cell_changes = pd.DataFrame({
        "Cycle": cycle, "RowKey": ch["RowKey"], "DeptCode": ch["DeptCode"].astype(int), "ProjectCode": ch["RowKey"].map(proj),
        "Month": pd.to_datetime(ch["Month"]), "Cell Planning MH": ch["P0"], "Discipline Value": ch["D"],
        "Cell Discipline MH": ch["Final"], "Change MH": ch["Change"], "Cell Implemented MH": ch["I0"],
        "Cell Status": ch["CellStatus"].where(ch["Changed"], "Blank retained"), "Blank Note": ch["AmbiguousReason"]})

    dim_dept = pd.DataFrame({"DeptCode": [int(d) for d in DEPARTMENTS], "Department": list(DEPARTMENTS.values())})
    dim_dept["Division"] = [int(DIVISION.get(str(d), d)) for d in dim_dept["DeptCode"]]
    dim_dept["Service"] = [SERVICE.get(str(d), "Engineering") for d in dim_dept["DeptCode"]]
    dim_dept["Dept Label"] = dim_dept["DeptCode"].astype(str) + " - " + dim_dept["Department"]
    capx = cap.copy()
    capx["DeptCode"] = capx["DeptCode"].astype(int)
    dim_dept = dim_dept.merge(capx.drop(columns=["HeadcountSource"]), on="DeptCode", how="left")
    if fb is not None:
        f = fb.assign(DeptCode=fb["DeptCode"].astype(int))[["DeptCode", "Feedback_Status", "Feedback_Processed",
                                                             "Tracker_Sent", "Tracker_Received"]]
        f = f.rename(columns={"Feedback_Status": "Feedback Status", "Feedback_Processed": "Feedback Processed",
                              "Tracker_Sent": "Sent To Dept", "Tracker_Received": "Comments Received"})
        f["Feedback Processed"] = f["Feedback Processed"].map({True: "Yes", False: "No"})
        dim_dept = dim_dept.merge(f, on="DeptCode", how="left")

    names = r.assign(N=r["D_Name"].fillna(r["P_Name"]).fillna(r["I_Name"])).groupby("ProjectFull")["N"].agg(lambda s: max(s, key=len))
    dim_proj = pd.DataFrame({"ProjectCode": names.index, "Project Name": names.values})
    dim_proj["Base Code"] = dim_proj["ProjectCode"].str.split("-").str[0]
    dim_proj["Client"] = dim_proj["Project Name"].map(_client)
    if "Master_Name" in ps:
        m = ps.set_index("ProjectFull")[["Master_Name", "Master_Client", "Master_Category", "Master_Flag", "Title_Variants"]]
        dim_proj = dim_proj.merge(m.rename(columns={"Master_Name": "Master Name", "Master_Client": "Master Client",
                                                    "Master_Category": "Master Category", "Master_Flag": "Master Flag",
                                                    "Title_Variants": "Title Variants"}),
                                  left_on="ProjectCode", right_index=True, how="left")
        dim_proj["Client"] = dim_proj["Master Client"].fillna(dim_proj["Client"])
    dim_proj["Project Label"] = dim_proj["ProjectCode"] + "  " + dim_proj["Project Name"].str.replace(r"^\s*[\d\-]+\s*", "", regex=True).str[:70]
    dim_proj = dim_proj.merge(sev[["Severity", "Severity_Points", "Depts_Concerned", "Departments_Raising"]]
                              .rename(columns={"Severity_Points": "Severity Points", "Depts_Concerned": "Depts Concerned",
                                               "Departments_Raising": "Departments Raising"}),
                              left_on="ProjectCode", right_index=True, how="left")
    dim_proj["Severity"] = dim_proj["Severity"].fillna("No Concern")
    dim_proj["Severity Order"] = dim_proj["Severity"].map({"Critical": 1, "High": 2, "Medium": 3, "Low": 4, "No Concern": 5})

    dim_cat = pd.DataFrame({"Category": CATEGORY_ORDER, "Category Order": range(1, len(CATEGORY_ORDER) + 1),
                            "Parent Category": ["Running", "Running", "Loan To", "Under Signature", "Proposal", "Potential Opportunity"],
                            "Firm Workload": ["Yes" if x in FIRM_CATEGORIES else "No" for x in CATEGORY_ORDER],
                            "Color": [CATEGORY_COLORS[x] for x in CATEGORY_ORDER]})
    dim_scen = pd.DataFrame(SCENARIOS, columns=["Scenario", "Scenario Order", "Color", "Meaning"])
    months = sorted(pd.to_datetime(c["Month"].unique()))
    dim_date = pd.DataFrame({"Month": months})
    dim_date["Year"] = dim_date["Month"].dt.year
    dim_date["Quarter"] = "Q" + dim_date["Month"].dt.quarter.astype(str) + "-" + dim_date["Year"].astype(str)
    dim_date["Month Label"] = dim_date["Month"].dt.strftime("%b-%y")
    dim_date["Month Index"] = range(1, len(dim_date) + 1)
    dim_date["Period"] = ["Months %d-%d" % (6 * ((i - 1) // 6) + 1, 6 * ((i - 1) // 6) + 6) for i in dim_date["Month Index"]]
    dq = pd.DataFrame(issues)[["Severity", "Area", "Scenario", "File", "Issue", "Detail", "Impact_MH"]]
    dq["Detail"] = dq["Detail"].astype(str).str[:500]
    dq.insert(0, "Cycle", cycle)

    for t in (items, bridge, cell_changes):
        t.insert(2, "CycleRowKey", t["Cycle"] + "|" + t["RowKey"])
    # whole numbers stay whole in the CSV (no "62.0"), so Power Query's Int64 columns load cleanly
    for t, cols in ((items, ["Feedback Row", "Project Severity Points", "Project Severity Order"]), (dim_proj, ["Severity Points", "Depts Concerned", "Severity Order"]),
                    (dim_dept, ["Managers", "LoanTo_Unpaid", "Permanent", "Temporary", "Total", "Effective_Heads"])):
        for col in cols:
            if col in t:
                t[col] = pd.to_numeric(t[col], errors="coerce").round().astype("Int64")
    dim_dept.insert(len(dim_dept.columns), "As Of", dt.date.today())
    dim_dept.insert(len(dim_dept.columns), "Cycle", cycle)   # headcount and feedback status belong to this cycle
    tables = {"FactWorkload": fact, "FactCommentItems": items, "FactCellChanges": cell_changes, "BridgeConcernType": bridge,
              "DimDepartment": dim_dept, "DimProject": dim_proj, "DimCategory": dim_cat, "DimScenario": dim_scen,
              "DimDate": dim_date, "DataQuality": dq}
    if inventory is not None:
        tables["FileInventory"] = inventory.assign(Cycle=cycle)
    # always written (headers only when there is none), so a re-run never leaves the previous run's transfers behind
    pt_cols = ["Dept", "Reduced Project", "Reduced Phase", "Reduced Title", "Reduction MH", "Added Project", "Added Phase",
               "Added Title", "Added MH", "Match Basis", "Possible Offset MH", "Net of the Pair MH"]
    tables["PossibleTransfers"] = (possible if possible is not None and len(possible) else pd.DataFrame(columns=pt_cols)).assign(Cycle=cycle)
    for name, t in tables.items():
        t.to_csv(folder / f"{name}.csv", index=False, encoding="utf-8-sig", date_format="%Y-%m-%d")
    wb = Workbook()
    wb.remove(wb.active)
    for name, t in tables.items():
        ws = wb.create_sheet(name[:31])
        ws.append(list(t.columns))
        for rec in t.itertuples(index=False):
            ws.append([None if (v is pd.NA or (isinstance(v, float) and pd.isna(v)))
                       else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in rec])
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True)
        if len(t):
            from openpyxl.utils import get_column_letter
            ref = f"A1:{get_column_letter(len(t.columns))}{len(t) + 1}"
            tab = Table(displayName=name, ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(tab)
        for col_cells in ws.iter_cols(min_row=1, max_row=1):
            ws.column_dimensions[col_cells[0].column_letter].width = 16
        if "Month" in t.columns:
            mi = list(t.columns).index("Month") + 1
            for row in ws.iter_rows(min_row=2, min_col=mi, max_col=mi):
                row[0].number_format = "mmm-yy"
    wb.save(folder / "Workload_Comments_PowerBI_Model.xlsx")
    for name in ("queries.pq", "measures.dax", "ENPPI_Workload_Theme.json", "PowerBI_Build_Guide.md"):
        if (PACKAGE / name).exists():
            shutil.copy2(PACKAGE / name, folder / name)
    return tables
