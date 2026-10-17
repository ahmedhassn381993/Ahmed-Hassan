"""Write the management workbook. Summary sheets are formula-driven (SUMIFS/COUNTIFS over the
register sheets) so totals stay traceable; registers hold the engine's row/cell-level results."""
from __future__ import annotations

import datetime as dt

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from analysis import CATEGORY_ORDER, DEPARTMENTS, FIRM_CATEGORIES, HOURS_PER_MAN_MONTH, PERMANENT_EFFECTIVENESS

FONT = "Arial"
NAVY, INK, MUTED, GRID = "1F3864", "0B0B0B", "52514E", "D9D9D9"
SCEN = {"Planning Input": "2A78D6", "Discipline Comment": "EB6834", "System Implementation": "1BAF7A"}
CAT_FILL = {"Overall": "1F3864", "Running": "C8E6C9", "Internal / Unforeseen": "B2DFDB", "Loan To": "CFD8DC",
            "Under Signature": "E1BEE7", "Proposal": "FFE0B2", "Potential Opportunity": "FFF9C4"}
SEV_FILL = {"Critical": "C00000", "High": "F4B183", "Medium": "FFE699", "Low": "E2EFDA", "No Concern": "F2F2F2"}
MH = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
MON = "mmm-yy"
THIN = Side(style="thin", color=GRID)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _f(bold=False, color=INK, size=9, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _fill(hex6):
    return PatternFill("solid", fgColor=hex6)


def _title(ws, text, sub=None, width_cols=12):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _f(color=MUTED, italic=True)
    ws.sheet_view.showGridLines = False


def _header(ws, row, col, headers, fill=NAVY, height=30):
    for j, h in enumerate(headers):
        c = ws.cell(row, col + j, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        c.fill = _fill(fill)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[row].height = height


def _write_table(ws, df: pd.DataFrame, row: int, col: int = 1, fmts: dict | None = None,
                 widths: dict | None = None, header_fill=NAVY, autofilter=True) -> tuple[int, int]:
    fmts, widths = fmts or {}, widths or {}
    _header(ws, row, col, list(df.columns), fill=header_fill)
    for i, rec in enumerate(df.itertuples(index=False), start=row + 1):
        for j, v in enumerate(rec):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            elif isinstance(v, (np.bool_, bool)):
                v = "Yes" if v else "No"
            c = ws.cell(i, col + j, v)
            c.font = _f()
            c.border = BOX
            name = df.columns[j]
            if name in fmts:
                c.number_format = fmts[name]
    last = row + len(df)
    for j, name in enumerate(df.columns):
        ws.column_dimensions[L(col + j)].width = widths.get(name, max(10, min(45, len(str(name)) + 4)))
    if autofilter and len(df):
        ws.auto_filter.ref = f"{L(col)}{row}:{L(col + len(df.columns) - 1)}{last}"
    return row + 1, last


def _note(ws, cell, text):
    ws[cell].comment = Comment(text, "Comments Analysis Engine")


# ----------------------------------------------------------------------------- data sheets
RR_COLS = ["DeptCode", "Department", "SectionCode", "Phase", "Category", "ProjectFull", "ProjectBase", "Activity",
           "P_Exists", "D_Exists", "I_Exists", "P_Total", "Final_Total", "I_Total", "Gross_Increase",
           "Gross_Reduction", "Net_Change", "Changed_Cells", "Concern", "Primary_Type", "Concern_Types",
           "Extension_MH", "Ambiguous_Cells", "Ambiguous_MH", "Impl_Ratio", "Implementation_Status",
           "Unrequested_Gross", "Unrequested_Net", "RequestedPhase", "CodeCorrection", "Months_Changed",
           "D_File", "D_Row", "P_Row", "I_Row", "Implemented_Request_MH", "Gross_Change"]
RR = {c: L(i + 1) for i, c in enumerate(RR_COLS)}


def _row_register(wb, rows):
    ws = wb.create_sheet("Row Register")
    r = rows.copy()
    r["DeptCode"] = r["DeptCode"].astype(int)
    r["Department"] = r["DeptCode"].astype(str).map(DEPARTMENTS)
    r["Activity"] = r["D_Name"].fillna(r["P_Name"]).fillna(r["I_Name"])
    r["Gross_Change"] = r["Gross_Increase"] - r["Gross_Reduction"]
    r["Implemented_Request_MH"] = np.where(r["Concern"], r["Impl_Ratio"].fillna(0) * r["Gross_Change"], 0.0)
    r = r.sort_values(["DeptCode", "SectionCode", "Phase", "ProjectFull"])[RR_COLS]
    _title(ws, "Row Register - one row per Department / Sub-department / Phase / Project",
           "Engine output (data sheet). Summary sheets read this sheet with SUMIFS / COUNTIFS.")
    fm = {c: MH for c in ["P_Total", "Final_Total", "I_Total", "Gross_Increase", "Gross_Reduction", "Net_Change",
                          "Extension_MH", "Ambiguous_MH", "Unrequested_Gross", "Unrequested_Net",
                          "Implemented_Request_MH", "Gross_Change"]}
    fm["Impl_Ratio"] = PCT
    first, last = _write_table(ws, r, 4, fmts=fm, widths={"Activity": 50, "Concern_Types": 40, "Months_Changed": 30})
    ws.freeze_panes = "I5"
    return first, last


def _monthly_data(wb, mdf):
    ws = wb.create_sheet("Monthly Data")
    m = mdf.copy()
    m["DeptCode"] = m["DeptCode"].astype(int)
    m = m.sort_values(["DeptCode", "Category", "Month"])
    _title(ws, "Monthly Data - Department x Category x Month (three scenarios)", "Engine output (data sheet).")
    first, last = _write_table(ws, m, 4, fmts={"Month": MON, "Planning_MH": MH, "Discipline_Final_MH": MH, "Implemented_MH": MH})
    for rr in range(first, last + 1):
        ws.cell(rr, 3).value = dt.datetime.combine(ws.cell(rr, 3).value, dt.time())
        ws.cell(rr, 3).number_format = MON
    ws.freeze_panes = "A5"
    return first, last


def _capacity(wb, cap):
    ws = wb.create_sheet("Capacity")
    c = cap.copy()
    c["DeptCode"] = c["DeptCode"].astype(int)
    _title(ws, "Department Capacity (headcount basis)",
           f"Source: employee_count.xlsx, sheet 'Table for Emails'. Effective heads = ROUNDUP(Permanent x "
           f"{PERMANENT_EFFECTIVENESS}) + Temporary; capacity = effective heads x {HOURS_PER_MAN_MONTH} MH/month "
           "(same factors as the Forecast Comments PBIX).")
    cols = ["DeptCode", "HeadcountSource", "Managers", "LoanTo_Unpaid", "Permanent", "Temporary", "Total"]
    first, last = _write_table(ws, c[cols], 4, autofilter=False)
    _header(ws, 4, 8, ["Effective_Heads", "Capacity_MH_per_Month"])
    ws["K3"], ws["K4"] = "Permanent effectiveness", PERMANENT_EFFECTIVENESS
    ws["L3"], ws["L4"] = "MH per man-month", HOURS_PER_MAN_MONTH
    for cell in ("K3", "L3"):
        ws[cell].font = _f(bold=True)
    for cell in ("K4", "L4"):
        ws[cell].font = Font(name=FONT, color="0000FF", size=9)
        ws[cell].fill = _fill("FFFF00")
    _note(ws, "K4", "Assumption from the Forecast Comments PBIX (Dept query: Perm_EFF = ROUNDUP(Permanent_No*0.9)).")
    _note(ws, "L4", "Assumption from the Forecast Comments PBIX (AvailMHR = EffTotal*Months*150).")
    for rr in range(first, last + 1):
        ws.cell(rr, 8, f"=ROUNDUP(E{rr}*$K$4,0)+F{rr}").font = _f()
        ws.cell(rr, 9, f"=H{rr}*$L$4").font = _f()
        ws.cell(rr, 9).number_format = MH
        for cc in (8, 9):
            ws.cell(rr, cc).border = BOX
    ws.column_dimensions["H"].width = ws.column_dimensions["I"].width = 16
    ws.column_dimensions["K"].width = ws.column_dimensions["L"].width = 22
    return first, last


# ----------------------------------------------------------------------------- summary sheets
def _dept_ranking(wb, ds, rr_last, md_last, cap_last):
    ws = wb["Department Ranking"]
    _title(ws, "Department Ranking - Planning Input vs Discipline Comment vs After Implementation",
           "Sorted by absolute net deviation. MH = man-hours over the 24-month window. Formulas read 'Row Register', "
           "'Monthly Data' and 'Capacity'.")
    heads = ["Dept", "Department", "Feedback status", "Planning Input MH", "Discipline-adjusted MH",
             "After Implementation MH", "Gross Increase", "Gross Reduction", "Net Deviation", "Deviation %",
             "|Net|", "Rank |Net|", "Concern Items", "Concern Projects", "Changed Cells", "New Projects",
             "Phase Transfers / Moves", "Schedule Extensions", "Blank/zero not applied MH", "Requested Change MH",
             "Implemented of Request MH", "Implementation Rate", "Planner Change w/o Comment MH",
             "Effective Heads", "Capacity MH / Month", "Avg Utilisation - Planning", "Avg Utilisation - Discipline",
             "Avg Utilisation - After Impl.", "Firm Utilisation - Discipline", "Rank by Concern Items"]
    _header(ws, 4, 1, heads, height=48)
    d = ds.assign(absnet=ds["Net_Deviation"].abs()).sort_values(["absnet", "Concern_Items"], ascending=False)
    R = "'Row Register'!"
    rng = lambda col: f"{R}${RR[col]}$5:${RR[col]}${rr_last}"
    first = 5
    last = first + len(d) - 1
    M = "'Monthly Data'!"
    for i, rec in enumerate(d.itertuples(), start=first):
        ws.cell(i, 1, int(rec.DeptCode))
        ws.cell(i, 2, rec.Department)
        ws.cell(i, 3, rec.Feedback_Status)
        crit = f"$A{i}"
        ws.cell(i, 4, f"=SUMIFS({rng('P_Total')},{rng('DeptCode')},{crit})")
        ws.cell(i, 5, f"=SUMIFS({rng('Final_Total')},{rng('DeptCode')},{crit})")
        ws.cell(i, 6, f"=SUMIFS({rng('I_Total')},{rng('DeptCode')},{crit})")
        ws.cell(i, 7, f"=SUMIFS({rng('Gross_Increase')},{rng('DeptCode')},{crit})")
        ws.cell(i, 8, f"=SUMIFS({rng('Gross_Reduction')},{rng('DeptCode')},{crit})")
        ws.cell(i, 9, f"=G{i}+H{i}")
        ws.cell(i, 10, f"=IF(D{i}=0,0,I{i}/D{i})")
        ws.cell(i, 11, f"=ABS(I{i})")
        ws.cell(i, 12, f'=IF(K{i}=0,"-",RANK(K{i},$K${first}:$K${last}))')
        ws.cell(i, 13, f'=COUNTIFS({rng("DeptCode")},{crit},{rng("Concern")},"Yes")')
        ws.cell(i, 14, int(rec.Concern_Projects))
        ws.cell(i, 15, f"=SUMIFS({rng('Changed_Cells')},{rng('DeptCode')},{crit})")
        ws.cell(i, 16, f'=COUNTIFS({rng("DeptCode")},{crit},{rng("Primary_Type")},"New Project*")')
        ws.cell(i, 17, f'=COUNTIFS({rng("DeptCode")},{crit},{rng("Concern_Types")},"*Phase Transfer*")'
                       f'+COUNTIFS({rng("DeptCode")},{crit},{rng("Concern_Types")},"*Phase Label*")')
        ws.cell(i, 18, f'=COUNTIFS({rng("DeptCode")},{crit},{rng("Concern_Types")},"*Schedule Extension*")')
        ws.cell(i, 19, f"=SUMIFS({rng('Ambiguous_MH')},{rng('DeptCode')},{crit})")
        ws.cell(i, 20, f"=G{i}-H{i}")
        ws.cell(i, 21, f"=SUMIFS({rng('Implemented_Request_MH')},{rng('DeptCode')},{crit})")
        ws.cell(i, 22, f'=IF(T{i}=0,"-",U{i}/T{i})')
        ws.cell(i, 23, f"=SUMIFS({rng('Unrequested_Gross')},{rng('DeptCode')},{crit})")
        ws.cell(i, 24, f"=IFERROR(INDEX(Capacity!$H$5:$H${cap_last},MATCH($A{i},Capacity!$A$5:$A${cap_last},0)),0)")
        ws.cell(i, 25, f"=IFERROR(INDEX(Capacity!$I$5:$I${cap_last},MATCH($A{i},Capacity!$A$5:$A${cap_last},0)),0)")
        ws.cell(i, 26, f'=IF(Y{i}=0,"-",D{i}/24/Y{i})')
        ws.cell(i, 27, f'=IF(Y{i}=0,"-",E{i}/24/Y{i})')
        ws.cell(i, 28, f'=IF(Y{i}=0,"-",F{i}/24/Y{i})')
        firm = "+".join(f'SUMIFS({M}$E$5:$E${md_last},{M}$A$5:$A${md_last},$A{i},{M}$B$5:$B${md_last},"{c}")'
                        for c in FIRM_CATEGORIES)
        ws.cell(i, 29, f'=IF(Y{i}=0,"-",({firm})/24/Y{i})')
        ws.cell(i, 30, f'=IF(M{i}=0,"-",RANK(M{i},$M${first}:$M${last}))')
        for j in range(1, 31):
            c = ws.cell(i, j)
            c.font, c.border = _f(), BOX
            c.number_format = PCT if j in (10, 22, 26, 27, 28, 29) else (MH if j in (4, 5, 6, 7, 8, 9, 11, 19, 20, 21, 23, 25) else "General")
    t = last + 1
    ws.cell(t, 1, "TOTAL")
    ws.cell(t, 2, "All departments")
    for j in (4, 5, 6, 7, 8, 9, 13, 14, 15, 16, 17, 18, 19, 20, 21, 23, 24, 25):
        ws.cell(t, j, f"=SUM({L(j)}{first}:{L(j)}{last})")
    ws.cell(t, 10, f"=IF(D{t}=0,0,I{t}/D{t})")
    ws.cell(t, 22, f'=IF(T{t}=0,"-",U{t}/T{t})')
    ws.cell(t, 26, f'=IF(Y{t}=0,"-",D{t}/24/Y{t})')
    ws.cell(t, 27, f'=IF(Y{t}=0,"-",E{t}/24/Y{t})')
    ws.cell(t, 28, f'=IF(Y{t}=0,"-",F{t}/24/Y{t})')
    firm_all = "+".join(f'SUMIFS({M}$E$5:$E${md_last},{M}$B$5:$B${md_last},"{c}")' for c in FIRM_CATEGORIES)
    ws.cell(t, 29, f'=IF(Y{t}=0,"-",({firm_all})/24/Y{t})')
    for j in range(1, 31):
        c = ws.cell(t, j)
        c.font, c.border, c.fill = _f(bold=True), BOX, _fill("DDEBF7")
        c.number_format = PCT if j in (10, 22, 26, 27, 28, 29) else (MH if j not in (1, 2, 3) else "General")
    widths = {1: 7, 2: 30, 3: 30}
    for j in range(1, 31):
        ws.column_dimensions[L(j)].width = widths.get(j, 13)
    ws.freeze_panes = "D5"
    lim = float(max(1.0, d["Net_Deviation"].abs().max()))
    ws.conditional_formatting.add(f"I{first}:I{last}", _diverging(lim))
    for col in ("Z", "AA", "AB", "AC"):
        ws.conditional_formatting.add(f"{col}{first}:{col}{last}", CellIsRule(operator="greaterThan", formula=["1"],
                                      font=Font(name=FONT, bold=True, color="C00000")))
    _note(ws, "N4", "Distinct project codes with at least one concern row (engine-computed distinct count).")
    _note(ws, "S4", "Baseline MH where feedback was blank but marked red/highlighted, or where the Copilot table "
                    "cannot distinguish blank from zero. Retained per rule 'blank = no revision'; verify with the department.")
    _note(ws, "W4", "Absolute MH changed in the after-implementation data where the department requested no change.")
    _note(ws, "AC4", "Running + Internal + Loan To + Under Signature (pipeline excluded), as the PBIX 'Projects Total%' measure.")
    ws.cell(t + 2, 1, "Utilisation = workload MH per month / (effective heads x 150 MH). Values above 100% are bold red. "
                      "Deviation colour: blue = reduction, orange = increase.").font = _f(color=MUTED, italic=True)
    return first, last, t


def _project_ranking(wb, ps, rr_last):
    ws = wb["Project Ranking"]
    _title(ws, "Project Ranking - departmental concern by project",
           "Severity points are formulas (see 'Method & Assumptions'). MH sums read 'Row Register'.")
    p = ps[ps["Concern_Rows"] > 0].copy()
    heads = ["Rank", "Project Code", "Base Code", "Project Name", "Severity", "Severity Points", "Depts Concerned",
             "Departments Raising", "Planning Input MH", "Discipline-adjusted MH", "After Implementation MH",
             "Gross Increase", "Gross Reduction", "Net Deviation", "Net %", "Months Affected", "Affected Months",
             "Affected Phases", "Concern Types", "New Allocation", "Schedule Extension", "Implementation Rate"]
    _header(ws, 4, 1, heads, height=40)
    R = "'Row Register'!"
    rng = lambda col: f"{R}${RR[col]}$5:${RR[col]}${rr_last}"
    first, last = 5, 4 + len(p)
    for i, rec in enumerate(p.itertuples(), start=first):
        ws.cell(i, 1, f"=RANK(F{i},$F${first}:$F${last})")
        ws.cell(i, 2, rec.ProjectFull)
        ws.cell(i, 3, rec.Project_Base)
        ws.cell(i, 4, rec.Project_Name[:120])
        ws.cell(i, 5, f'=IF(F{i}>=10,"Critical",IF(F{i}>=7,"High",IF(F{i}>=4,"Medium","Low")))')
        g = f"(L{i}-M{i})"
        ws.cell(i, 6, f"=IF({g}>=20000,5,IF({g}>=10000,4,IF({g}>=5000,3,IF({g}>=2000,2,IF({g}>=500,1,0)))))"
                      f"+IF(I{i}<=0,2,IF(ABS(N{i})/I{i}>=0.5,2,IF(ABS(N{i})/I{i}>=0.2,1,0)))"
                      f"+IF(G{i}>=6,3,IF(G{i}>=4,2,IF(G{i}>=2,1,0)))+IF(P{i}>=12,2,IF(P{i}>=6,1,0))"
                      f'+IF(T{i}="Yes",1,0)+IF(U{i}="Yes",1,0)')
        ws.cell(i, 7, int(rec.Depts_Concerned))
        ws.cell(i, 8, rec.Departments_Raising)
        crit = f"$B{i}"
        for j, col in ((9, "P_Total"), (10, "Final_Total"), (11, "I_Total"), (12, "Gross_Increase"), (13, "Gross_Reduction")):
            ws.cell(i, j, f"=SUMIFS({rng(col)},{rng('ProjectFull')},{crit})")
        ws.cell(i, 14, f"=L{i}+M{i}")
        ws.cell(i, 15, f'=IF(I{i}<=0,"New",N{i}/I{i})')
        ws.cell(i, 16, int(rec.Months_Affected))
        ws.cell(i, 17, rec.Affected_Months)
        ws.cell(i, 18, rec.Affected_Phases)
        ws.cell(i, 19, rec.Concern_Types)
        ws.cell(i, 20, "Yes" if rec.Has_New_Allocation else "No")
        ws.cell(i, 21, "Yes" if rec.Has_Extension else "No")
        ws.cell(i, 22, f'=IF((L{i}-M{i})=0,"-",SUMIFS({rng("Implemented_Request_MH")},{rng("ProjectFull")},{crit})/(L{i}-M{i}))')
        for j in range(1, 23):
            c = ws.cell(i, j)
            c.font, c.border = _f(), BOX
            c.number_format = PCT if j in (15, 22) else (MH if j in (9, 10, 11, 12, 13, 14) else "General")
    for sev, col in SEV_FILL.items():
        ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="equal", formula=[f'"{sev}"'], fill=_fill(col),
                                      font=Font(name=FONT, bold=True, color="FFFFFF" if sev == "Critical" else INK)))
    widths = {1: 6, 2: 12, 3: 10, 4: 48, 5: 10, 6: 9, 7: 9, 8: 30, 17: 18, 18: 26, 19: 46}
    for j in range(1, 23):
        ws.column_dimensions[L(j)].width = widths.get(j, 13)
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:V{last}"
    return first, last


def _diverging(lim: float) -> ColorScaleRule:
    """Blue = reduction, neutral grey at zero, orange = increase; symmetric so zero is always neutral."""
    return ColorScaleRule(start_type="num", start_value=-lim, start_color="9EC3EE", mid_type="num", mid_value=0,
                          mid_color="F7F7F7", end_type="num", end_value=lim, end_color="F3A47F")


def _monthly_impact(wb, months, md_last, cap_last, dev_lim):
    ws = wb["Monthly Workload Impact"]
    _title(ws, "Monthly Workload Impact - company and department",
           "All values are SUMIFS over 'Monthly Data'. Discipline-adjusted = Planning Input with populated, "
           "different feedback values applied.")
    M = "'Monthly Data'!"
    rng = lambda c: f"{M}${c}$5:${c}${md_last}"
    hr = 4
    ws.cell(hr, 1, "Measure").font = _f(bold=True)
    ws.cell(hr, 2, "Category / Scenario").font = _f(bold=True)
    for j, m in enumerate(months, start=3):
        c = ws.cell(hr, j, dt.datetime.combine(m, dt.time()))
        c.number_format = MON
    tot_col = 3 + len(months)
    ws.cell(hr, tot_col, "Total 24 m")
    _header(ws, hr, 1, [ws.cell(hr, j).value for j in range(1, tot_col + 1)])
    for j in range(3, tot_col):
        ws.cell(hr, j).number_format = MON
    row = hr + 1
    scen_cols = [("Planning Input", "D"), ("Discipline-adjusted", "E"), ("After Implementation", "F")]
    total_rows = {}
    for label, col in scen_cols:
        for cat in CATEGORY_ORDER:
            ws.cell(row, 1, label)
            ws.cell(row, 2, cat)
            for j in range(3, tot_col):
                ws.cell(row, j, f"=SUMIFS({rng(col)},{rng('C')},{L(j)}${hr},{rng('B')},$B{row})")
            ws.cell(row, tot_col, f"=SUM({L(3)}{row}:{L(tot_col - 1)}{row})")
            for j in range(1, tot_col + 1):
                c = ws.cell(row, j)
                c.font, c.border, c.number_format = _f(), BOX, MH
                c.fill = _fill(CAT_FILL[cat])
            row += 1
        ws.cell(row, 1, label)
        ws.cell(row, 2, "TOTAL")
        for j in range(3, tot_col + 1):
            ws.cell(row, j, f"=SUM({L(j)}{row - len(CATEGORY_ORDER)}:{L(j)}{row - 1})")
        for j in range(1, tot_col + 1):
            c = ws.cell(row, j)
            c.font, c.border, c.number_format = _f(bold=True, color="FFFFFF"), BOX, MH
            c.fill = _fill(NAVY)
        total_rows[label] = row
        row += 2
    p, d, i = total_rows["Planning Input"], total_rows["Discipline-adjusted"], total_rows["After Implementation"]
    lines = [("Deviation", "Discipline - Planning", lambda c: f"={c}{d}-{c}{p}"),
             ("Deviation", "After Impl. - Planning", lambda c: f"={c}{i}-{c}{p}"),
             ("Deviation", "After Impl. - Discipline", lambda c: f"={c}{i}-{c}{d}"),
             ("Capacity", "Company capacity MH (effective heads x 150)", lambda c: f"=SUM(Capacity!$I$5:$I${cap_last})"),
             ("Utilisation", "Planning Input", lambda c: f"=IF({c}$CAP=0,0,{c}{p}/{c}$CAP)"),
             ("Utilisation", "Discipline-adjusted", lambda c: f"=IF({c}$CAP=0,0,{c}{d}/{c}$CAP)"),
             ("Utilisation", "After Implementation", lambda c: f"=IF({c}$CAP=0,0,{c}{i}/{c}$CAP)"),
             ("Utilisation", "Firm (Running+Internal+Loan To+Under Signature) - Discipline",
              lambda c: f"=IF({c}$CAP=0,0,SUM({c}{d - 6}:{c}{d - 3})/{c}$CAP)")]
    cap_row = row + 3
    for meas, lab, fn in lines:
        ws.cell(row, 1, meas)
        ws.cell(row, 2, lab)
        for j in range(3, tot_col + 1):
            colL = L(j)
            ws.cell(row, j, fn(colL).replace("$CAP", str(cap_row)))
            c = ws.cell(row, j)
            c.font, c.border = _f(bold=meas != "Utilisation"), BOX
            c.number_format = PCT if meas == "Utilisation" else MH
        row += 1
    ws.cell(cap_row, tot_col, f"={L(tot_col - 1)}{cap_row}*24")
    util_rows = (cap_row + 1, cap_row + 4)
    ws.conditional_formatting.add(f"C{util_rows[0]}:{L(tot_col - 1)}{util_rows[1]}",
                                  CellIsRule(operator="greaterThan", formula=["1"], font=Font(name=FONT, bold=True, color="C00000")))
    # chart: company totals vs capacity (one axis: MH)
    ch = LineChart()
    ch.title = "Company workload by month - three scenarios vs capacity (MH)"
    ch.height, ch.width = 9, 26
    ch.y_axis.title, ch.x_axis.title = "Man-hours", None
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E7E6E6"))
    ch.legend.position = "b"
    for label, colr in ((p, SCEN["Planning Input"]), (d, SCEN["Discipline Comment"]), (i, SCEN["System Implementation"]), (cap_row, "52514E")):
        ref = Reference(ws, min_col=2, max_col=tot_col - 1, min_row=label)
        ch.add_data(ref, titles_from_data=True, from_rows=True)
        s = ch.series[-1]
        s.graphicalProperties.line.solidFill = colr
        s.graphicalProperties.line.width = 22000
        s.smooth = False
        if label == cap_row:
            s.graphicalProperties.line.dashStyle = "dash"
    ch.set_categories(Reference(ws, min_col=3, max_col=tot_col - 1, min_row=hr))
    ch.x_axis.number_format = "mmm-yy"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    # series titles come from column B ('TOTAL'); name them explicitly
    from openpyxl.chart.series import SeriesLabel
    for s, name in zip(ch.series, ["Planning Input", "Discipline-adjusted", "After Implementation", "Capacity"]):
        s.tx = SeriesLabel(v=name)
    ws.add_chart(ch, f"A{row + 2}")
    row += 22
    # department x month heatmaps
    depts = sorted(int(x) for x in DEPARTMENTS)
    for title, expr, fmt, scale in (
            ("Department deviation by month: Discipline-adjusted minus Planning Input (MH)",
             lambda dep_row, cl: f"=SUMIFS({rng('E')},{rng('A')},$A{dep_row},{rng('C')},{cl}${hr})-SUMIFS({rng('D')},{rng('A')},$A{dep_row},{rng('C')},{cl}${hr})",
             MH, "div"),
            ("Department utilisation by month: Discipline-adjusted MH / capacity",
             lambda dep_row, cl: f"=IFERROR(SUMIFS({rng('E')},{rng('A')},$A{dep_row},{rng('C')},{cl}${hr})/INDEX(Capacity!$I$5:$I${cap_last},MATCH($A{dep_row},Capacity!$A$5:$A${cap_last},0)),0)",
             PCT, "seq")):
        ws.cell(row, 1, title).font = Font(name=FONT, bold=True, size=11, color=NAVY)
        row += 1
        _header(ws, row, 1, ["Dept", "Department"] + [m.strftime("%b-%y") for m in months] + (["Total"] if fmt == MH else ["Peak"]))
        row += 1
        top = row
        for dep in depts:
            ws.cell(row, 1, dep)
            ws.cell(row, 2, DEPARTMENTS[str(dep)])
            for j in range(3, tot_col):
                ws.cell(row, j, expr(row, L(j)))
            ws.cell(row, tot_col, f"=SUM(C{row}:{L(tot_col - 1)}{row})" if fmt == MH else f"=MAX(C{row}:{L(tot_col - 1)}{row})")
            for j in range(1, tot_col + 1):
                c = ws.cell(row, j)
                c.font, c.border = _f(), BOX
                c.number_format = fmt if j > 2 else "General"
            row += 1
        area = f"C{top}:{L(tot_col - 1)}{row - 1}"
        if scale == "div":
            ws.conditional_formatting.add(area, _diverging(dev_lim))
        else:
            ws.conditional_formatting.add(area, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                               mid_type="num", mid_value=1, mid_color="FBD3BF",
                                                               end_type="num", end_value=2, end_color="D95926"))
            ws.conditional_formatting.add(area, CellIsRule(operator="greaterThan", formula=["1"], font=Font(name=FONT, bold=True, color=INK)))
        row += 2
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 34
    for j in range(3, tot_col + 1):
        ws.column_dimensions[L(j)].width = 9.5
    ws.freeze_panes = "C5"
    return total_rows, cap_row


def _deviation_register(wb, rows, ps):
    ws = wb["Deviation Register"]
    _title(ws, "Detailed Deviation Register - every concern, phase move, retained row and unresolved blank",
           "One row per Department / Sub-department / Phase / Project. Net % = Net / Planning Input.")
    sev = ps.set_index("ProjectFull")["Severity"]
    r = rows[rows["Concern"] | rows["Primary_Type"].isin(["Phase Label Change Only", "Description Change Only",
                                                          "Not in Feedback (retained)"]) | (rows["Ambiguous_Cells"] > 0)].copy()
    r["Activity"] = r["D_Name"].fillna(r["P_Name"])
    r["Project Severity"] = r["ProjectFull"].map(sev)
    r["Dept"] = r["DeptCode"].astype(int)
    r = r.assign(absnet=r["Net_Change"].abs()).sort_values(["Dept", "absnet"], ascending=[True, False])
    out = pd.DataFrame({
        "Dept": r["Dept"], "Department": r["DeptCode"].map(DEPARTMENTS), "Sub-dept": r["SectionCode"],
        "Phase": r["Phase"], "Category": r["Category"], "Project Code": r["ProjectFull"], "Activity (department wording)": r["Activity"],
        "Planning Input MH": r["P_Total"], "Discipline-adjusted MH": r["Final_Total"], "After Impl. MH": r["I_Total"],
        "Gross Increase": r["Gross_Increase"], "Gross Reduction": r["Gross_Reduction"], "Net Deviation": r["Net_Change"],
        "Net %": None, "Months Changed": r["Months_Changed"], "Concern Types": r["Concern_Types"],
        "Concern": r["Concern"], "Project Severity": r["Project Severity"], "Implementation Status": r["Implementation_Status"],
        "Implementation Rate": r["Impl_Ratio"], "Requested Phase": r["RequestedPhase"], "Code Correction": r["CodeCorrection"],
        "Blank/zero not applied MH": r["Ambiguous_MH"], "Feedback File": r["D_File"], "Feedback Row": r["D_Row"]})
    fm = {c: MH for c in ["Planning Input MH", "Discipline-adjusted MH", "After Impl. MH", "Gross Increase",
                          "Gross Reduction", "Net Deviation", "Blank/zero not applied MH"]}
    fm.update({"Net %": PCT, "Implementation Rate": PCT})
    first, last = _write_table(ws, out, 4, fmts=fm, widths={"Activity (department wording)": 50, "Concern Types": 40,
                                                              "Months Changed": 28, "Department": 26, "Implementation Status": 28})
    for i in range(first, last + 1):
        ws.cell(i, 14, f'=IF(H{i}<=0,"New",M{i}/H{i})').number_format = PCT
    for sevname, col in SEV_FILL.items():
        ws.conditional_formatting.add(f"R{first}:R{last}", CellIsRule(operator="equal", formula=[f'"{sevname}"'], fill=_fill(col),
                                      font=Font(name=FONT, color="FFFFFF" if sevname == "Critical" else INK)))
    ws.freeze_panes = "H5"
    return len(out)


def _cell_changes(wb, cells, rows):
    ws = wb["Cell Changes"]
    _title(ws, "Cell Changes - month-level changes and unresolved blanks",
           "Discipline value '(blank)' = no revision; the baseline is retained. Final = value used for the discipline-adjusted forecast.")
    k = ["DeptCode", "SectionCode", "Phase", "ProjectKey"]
    c = cells[cells["Changed"] | cells["AmbiguousBlank"]].merge(rows[k + ["ProjectFull", "D_Name", "P_Name", "Category"]], on=k)
    c = c.sort_values(["DeptCode", "ProjectFull", "Phase", "Month"])
    out = pd.DataFrame({
        "Dept": c["DeptCode"].astype(int), "Sub-dept": c["SectionCode"], "Phase": c["Phase"], "Project Code": c["ProjectFull"],
        "Activity": c["D_Name"].fillna(c["P_Name"]), "Month": [dt.datetime.combine(m, dt.time()) for m in c["Month"]],
        "Planning Input": c["P0"], "Discipline Value": c["D"].where(c["D"].notna(), "(blank)"),
        "Final (rule applied)": c["Final"], "Deviation": c["Change"], "After Implementation": c["I0"],
        "Implementation Gap (After - Final)": c["ImplGap"], "Cell Status": np.where(c["Changed"], c["CellStatus"], "Blank retained"),
        "Department Marking": np.where(c["D_Red"].fillna(False).astype(bool), "red font", "") + np.where(c["D_Fill"].fillna("").ne(""), " / fill " + c["D_Fill"].fillna(""), ""),
        "Blank Note": c["AmbiguousReason"]})
    fm = {"Month": MON, "Planning Input": MH, "Discipline Value": MH, "Final (rule applied)": MH, "Deviation": MH,
          "After Implementation": MH, "Implementation Gap (After - Final)": MH}
    first, last = _write_table(ws, out, 4, fmts=fm, widths={"Activity": 44, "Blank Note": 60, "Cell Status": 22})
    ws.conditional_formatting.add(f"H{first}:H{last}", FormulaRule(formula=[f'AND(ISNUMBER(H{first}),H{first}<>G{first})'],
                                  fill=_fill("FCE4D6"), font=Font(name=FONT, color="C55A11", bold=True)))
    ws.freeze_panes = "F5"
    return len(out)


def _new_projects(wb, rows):
    ws = wb["New Projects Register"]
    _title(ws, "New Projects Register - allocations that exist only in department feedback",
           "'Known in other departments' = the project code exists in another department's baseline.")
    r = rows[rows["Concern_Types"].str.contains("New Project|New Phase", regex=True)].copy()
    raised = r.groupby("ProjectFull")["DeptCode"].nunique()
    r = r.sort_values(["Final_Total"], ascending=False)
    out = pd.DataFrame({
        "Project Code": r["ProjectFull"], "Activity (department wording)": r["D_Name"], "Dept": r["DeptCode"].astype(int),
        "Department": r["DeptCode"].map(DEPARTMENTS), "Phase": r["Phase"], "Type": r["Primary_Type"],
        "Requested MH (24 m)": r["Final_Total"], "First Month": r["Final_First"], "Last Month": r["Final_Last"],
        "Departments Requesting Same Code": r["ProjectFull"].map(raised),
        "Known in Other Depts' Baseline": r["Known_In_Other_Depts"], "In After-Implementation Data": r["I_Exists"],
        "After Implementation MH": r["I_Total"], "Implementation Status": r["Implementation_Status"],
        "Feedback File": r["D_File"], "Row": r["D_Row"]})
    out["First Month"] = [dt.datetime.combine(m, dt.time()) if m is not None and not pd.isna(m) else None for m in out["First Month"]]
    out["Last Month"] = [dt.datetime.combine(m, dt.time()) if m is not None and not pd.isna(m) else None for m in out["Last Month"]]
    _write_table(ws, out, 4, fmts={"Requested MH (24 m)": MH, "After Implementation MH": MH, "First Month": MON, "Last Month": MON},
                 widths={"Activity (department wording)": 55, "Department": 26, "Implementation Status": 34})
    ws.freeze_panes = "C5"
    return len(out)


def _phase_and_schedule(wb, rows, transfers, issues, possible=None):
    ws = wb["Phase Transfers & Rescheduling"]
    _title(ws, "Phase Transfers, Rescheduling and Double-Counting Risks",
           "Phase moves are paired with the baseline row and kept in the baseline phase until Corporate Planning approves the category change.")
    row = 4
    ws.cell(row, 1, "A. Projects moved to another phase in the feedback").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    r = rows[rows["RequestedPhase"] != ""].copy()
    a = pd.DataFrame({"Dept": r["DeptCode"].astype(int), "Project Code": r["ProjectFull"], "Activity": r["D_Name"],
                      "Baseline Phase": r["Phase"], "Phase in Feedback": r["RequestedPhase"], "Planning MH": r["P_Total"],
                      "Discipline-adjusted MH": r["Final_Total"], "Net Change": r["Net_Change"], "Classification": r["Primary_Type"],
                      "MH that would be double counted if not paired": r["Final_Total"]})
    _, row = _write_table(ws, a, row + 1, fmts={c: MH for c in ["Planning MH", "Discipline-adjusted MH", "Net Change",
                                                              "MH that would be double counted if not paired"]},
                          widths={"Activity": 50}, autofilter=False)
    row += 2
    ws.cell(row, 1, "B. Reduction in one phase with increase in another (same department, same base project)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    if len(transfers):
        b = transfers.rename(columns={"pos": "Increase", "neg": "Reduction", "phases_up": "Phases Up", "phases_down": "Phases Down"})
        _, row = _write_table(ws, b, row + 1, autofilter=False)
    else:
        row += 1
        ws.cell(row, 1, "None found: no department reduced a project in one phase while increasing it in another.").font = _f(italic=True, color=MUTED)
    row += 2
    ws.cell(row, 1, "B2. Possible transfers between project codes or phases (verify with the department)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    if possible is not None and len(possible):
        _, row = _write_table(ws, possible, row + 1, fmts={c: MH for c in ["Reduction MH", "Added MH", "Possible Offset MH", "Net of the Pair MH"]},
                              widths={"Reduced Title": 40, "Added Title": 40, "Match Basis": 30}, autofilter=False)
        row += 1
        ws.cell(row, 1, "Gross figures keep both rows; if confirmed, the offset is a transfer, not new workload.").font = _f(italic=True, color=MUTED)
    else:
        row += 1
        ws.cell(row, 1, "None found.").font = _f(italic=True, color=MUTED)
    row += 2
    ws.cell(row, 1, "C. Schedule changes (extension, earlier start, earlier finish, later start)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    s = rows[rows["Concern_Types"].str.contains("Schedule|Earlier Start|Rescheduled", regex=True)].copy()
    s = s.sort_values("Extension_MH", ascending=False)
    fm = lambda x: x.strftime("%b-%y") if x is not None and not pd.isna(x) else ""
    c = pd.DataFrame({"Dept": s["DeptCode"].astype(int), "Phase": s["Phase"], "Project Code": s["ProjectFull"], "Activity": s["D_Name"],
                      "Baseline Start": s["P_First"].map(fm), "Baseline Finish": s["P_Last"].map(fm),
                      "Requested Start": s["Final_First"].map(fm), "Requested Finish": s["Final_Last"].map(fm),
                      "MH Beyond Baseline Finish": s["Extension_MH"], "MH Before Baseline Start": s["EarlyStart_MH"],
                      "Net Change": s["Net_Change"], "Types": s["Concern_Types"], "Implementation Status": s["Implementation_Status"]})
    _, row = _write_table(ws, c, row + 1, fmts={"MH Beyond Baseline Finish": MH, "MH Before Baseline Start": MH, "Net Change": MH},
                          widths={"Activity": 46, "Types": 40, "Implementation Status": 28}, autofilter=False)
    row += 2
    ws.cell(row, 1, "D. Double-counting risk register").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    risk = [i for i in issues if i["Area"] in ("Double counting", "Copilot model", "Phase transfer", "Possible transfer",
                                               "Duplicate project row", "Project code", "Possible copy-paste")]
    dd = pd.DataFrame(risk)[["Severity", "Area", "Scenario", "File", "Issue", "Detail", "Impact_MH"]] if risk else pd.DataFrame()
    if len(dd):
        _write_table(ws, dd, row + 1, fmts={"Impact_MH": MH}, widths={"Issue": 50, "Detail": 90}, autofilter=False)
    ws.column_dimensions["A"].width = 10


def _implementation(wb, rows, rr_last):
    ws = wb["Implementation Tracking"]
    _title(ws, "Implementation Tracking - were discipline comments implemented by Planning?",
           "Compares each requested change with the after-implementation extract. Counts and MH are formulas over 'Row Register'.")
    statuses = ["Implemented", "Partially implemented", "Not implemented", "Not implemented (row absent after implementation)",
                "Planner change without comment", "Added by Planning without comment"]
    heads = ["Dept", "Department", "Concern Items"] + statuses + ["Requested Change MH", "Implemented of Request MH",
                                                                  "Implementation Rate", "Planner Change w/o Comment MH"]
    _header(ws, 4, 1, heads, height=48)
    R = "'Row Register'!"
    rng = lambda col: f"{R}${RR[col]}$5:${RR[col]}${rr_last}"
    first = 5
    depts = sorted(int(x) for x in DEPARTMENTS)
    for i, dep in enumerate(depts, start=first):
        ws.cell(i, 1, dep)
        ws.cell(i, 2, DEPARTMENTS[str(dep)])
        ws.cell(i, 3, f'=COUNTIFS({rng("DeptCode")},$A{i},{rng("Concern")},"Yes")')
        for j, st in enumerate(statuses, start=4):
            ws.cell(i, j, f'=COUNTIFS({rng("DeptCode")},$A{i},{rng("Implementation_Status")},"{st}")')
        ws.cell(i, 10, f'=SUMIFS({rng("Gross_Change")},{rng("DeptCode")},$A{i},{rng("Concern")},"Yes")')
        ws.cell(i, 11, f'=SUMIFS({rng("Implemented_Request_MH")},{rng("DeptCode")},$A{i})')
        ws.cell(i, 12, f'=IF(J{i}=0,"-",K{i}/J{i})')
        ws.cell(i, 13, f'=SUMIFS({rng("Unrequested_Gross")},{rng("DeptCode")},$A{i})')
        for j in range(1, 14):
            c = ws.cell(i, j)
            c.font, c.border = _f(), BOX
            c.number_format = PCT if j == 12 else (MH if j in (10, 11, 13) else "General")
    last = first + len(depts) - 1
    t = last + 1
    ws.cell(t, 1, "TOTAL")
    for j in range(3, 12):
        ws.cell(t, j, f"=SUM({L(j)}{first}:{L(j)}{last})")
    ws.cell(t, 12, f'=IF(J{t}=0,"-",K{t}/J{t})')
    ws.cell(t, 13, f"=SUM(M{first}:M{last})")
    for j in range(1, 14):
        c = ws.cell(t, j)
        c.font, c.border, c.fill = _f(bold=True), BOX, _fill("DDEBF7")
        c.number_format = PCT if j == 12 else (MH if j in (10, 11, 13) else "General")
    ws.conditional_formatting.add(f"L{first}:L{last}", ColorScaleRule(start_type="num", start_value=0, start_color="FBD3BF",
                                  end_type="num", end_value=1, end_color="FFFFFF"))
    for j, w in enumerate([7, 30, 10, 13, 13, 13, 16, 15, 15, 14, 14, 13, 15], start=1):
        ws.column_dimensions[L(j)].width = w
    row = t + 3
    ws.cell(row, 1, "Largest requests not implemented (top 40 by requested MH)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    r = rows[rows["Concern"] & rows["Implementation_Status"].str.startswith("Not implemented")].copy()
    r["Req"] = r["Gross_Increase"] - r["Gross_Reduction"]
    r = r.sort_values("Req", ascending=False).head(40)
    out = pd.DataFrame({"Dept": r["DeptCode"].astype(int), "Phase": r["Phase"], "Project Code": r["ProjectFull"], "Activity": r["D_Name"],
                        "Planning MH": r["P_Total"], "Requested MH": r["Final_Total"], "After Impl. MH": r["I_Total"],
                        "Net Requested": r["Net_Change"], "Status": r["Implementation_Status"], "Types": r["Concern_Types"]})
    _write_table(ws, out, row + 1, fmts={c: MH for c in ["Planning MH", "Requested MH", "After Impl. MH", "Net Requested"]},
                 widths={}, autofilter=False)
    ws.column_dimensions["D"].width = 13
    ws.freeze_panes = "C5"


def _checklist(wb, chk):
    ws = wb["Comments Checklist"]
    _title(ws, "Critical Projects - Workload Forecast Comments Checklist (auto-filled)",
           "Same layout as the planners' checklist. 'Not implemented' items need the planner to confirm "
           "'Comments received but not applicable' or implement.")
    depts = [c for c in chk.columns if c not in ("Project No.", "Project Title")]
    _header(ws, 4, 1, ["Sr.", "Project No.", "Project Title"] + [f"{d}\n{DEPARTMENTS.get(d, 'Piping Division (2410-2450)')}" for d in depts], height=80)
    colors = {"Comments received & implemented": "C6EFCE", "Comments received & partially implemented": "FFEB9C",
              "No discipline comments": "F2F2F2"}
    for i, rec in enumerate(chk.itertuples(index=False), start=5):
        ws.cell(i, 1, i - 4)
        ws.cell(i, 2, str(rec[0]))
        ws.cell(i, 3, rec[1])
        for j, v in enumerate(rec[2:], start=4):
            c = ws.cell(i, j, v)
            c.fill = _fill(colors.get(v, "FFC7CE"))
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for j in range(1, 4 + len(depts)):
            ws.cell(i, j).font, ws.cell(i, j).border = _f(size=8), BOX
        ws.row_dimensions[i].height = 42
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 40
    for j in range(4, 4 + len(depts)):
        ws.column_dimensions[L(j)].width = 14
    lr = 6 + len(chk)
    for k, (label, col) in enumerate([("Comments received & implemented", "C6EFCE"), ("Comments received & partially implemented", "FFEB9C"),
                                      ("Comments received - not implemented (planner to confirm if not applicable)", "FFC7CE"),
                                      ("No discipline comments", "F2F2F2")]):
        ws.cell(lr + k, 3, label).fill = _fill(col)
        ws.cell(lr + k, 3).font = _f()
    ws.freeze_panes = "D5"


def _dq(wb, inventory, issues):
    ws = wb["Data Quality Log"]
    _title(ws, "Data Quality, Missing Files and Corrections Log",
           "Every correction the engine made to reach the figures in this workbook, and every gap that needs action.")
    ws.cell(4, 1, "A. File inventory (26 departments x 3 scenarios)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    _, last = _write_table(ws, inventory, 5, autofilter=False, widths={"Action Needed": 60, "Department": 30})
    row = last + 3
    ws.cell(row, 1, "B. Issues and corrections").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    iss = pd.DataFrame(issues)[["Severity", "Area", "Scenario", "File", "Issue", "Detail", "Impact_MH"]]
    order = {"High": 0, "Medium": 1, "Warning": 1, "Low": 2, "Info": 3}
    iss = iss.assign(o=iss["Severity"].map(order).fillna(4)).sort_values(["o", "Area"]).drop(columns="o")
    _write_table(ws, iss, row + 1, fmts={"Impact_MH": MH}, widths={"Issue": 55, "Detail": 90, "Scenario": 20}, autofilter=False)
    ws.freeze_panes = "A6"


def _method(wb, meta):
    ws = wb["Method & Assumptions"]
    _title(ws, "Method, Rules and Assumptions")
    lines = [
        ("Sources", meta["sources"]),
        ("Comparison window", meta["window"]),
        ("Match key", "Department + Sub-department + Phase + full Project code (+ occurrence when a code repeats) + Month. "
                      "Full codes are taken from the activity text because the template's =LEFT(B,7) formula truncates "
                      "8-digit and suffixed codes (40061420 -> 4006142; 5020250-7 -> 5020250)."),
        ("Rule 1 - blank feedback", "Final = Planning Input; deviation 0; no concern. Blank is never treated as zero."),
        ("Rule 2 - equal value", "Final = Planning Input; deviation 0; no concern."),
        ("Rule 3 - populated & different", "Final = feedback value; deviation = feedback - planning; concern = Yes."),
        ("Rule 5 - new rows", "Rows only in feedback are added (New Project, or New Phase Allocation when the base project already "
                              "exists in that department's baseline)."),
        ("Rule 6 - missing rows", "Baseline rows missing from feedback are retained (not removed)."),
        ("Rule 7 - phase moves", "Same project found under another phase in the feedback is paired with its baseline row; the MH "
                                 "comparison is like-for-like and the baseline phase is kept until Corporate Planning approves the "
                                 "category change. Identical values = 'Phase Label Change Only' (usually a missing phase heading)."),
        ("Rule 8 - gross vs net", "Gross increase, gross reduction, net and any offsetting phase redistribution are reported separately."),
        ("Code corrections", "A feedback row with no baseline match is paired with a baseline row missing from feedback in the same "
                             "department/phase when titles match >= 80%."),
        ("Unresolved blanks", "Departments without uploaded -c files come from the Copilot table, which stored blank feedback as 0. "
                              "Those zeros are treated as blank (baseline retained) and listed as 'Blank/zero not applied'. "
                              "In the 7 files read directly, a blank cell the department marked red/highlighted is listed the same way."),
        ("Implementation", "Per changed cell: Implemented if After = Final; Not implemented if After = Planning; otherwise Partially "
                           "implemented (share = movement toward the request). Row rate is weighted by requested MH: >= 95% Implemented, "
                           ">= 5% Partially implemented, else Not implemented."),
        ("Severity points", "Gross change (Inc + |Red|): >=20k 5, >=10k 4, >=5k 3, >=2k 2, >=500 1 | Net % of planning: new or >=50% 2, "
                            ">=20% 1 | Departments raising: >=6 3, >=4 2, >=2 1 | Months affected: >=12 2, >=6 1 | +1 new allocation | "
                            "+1 schedule extension. Critical >= 10, High 7-9, Medium 4-6, Low 1-3."),
        ("Capacity", f"Effective heads = ROUNDUP(Permanent x {PERMANENT_EFFECTIVENESS}) + Temporary; capacity = heads x "
                     f"{HOURS_PER_MAN_MONTH} MH/month (factors from the Forecast Comments PBIX). Managers and Loan-to/unpaid staff excluded. "
                     "Headcount date of 'Table for Emails' is not stated in the file - confirm it matches this cycle."),
        ("Firm workload", "Running + Internal / Unforeseen + Loan To + Under Signature (pipeline categories excluded)."),
        ("Category rule", "Internal / Unforeseen = project 1100100 under Running; it remains part of Running totals."),
        ("Limitations", meta["limitations"]),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(i, 1, k).font = _f(bold=True)
        ws.cell(i, 2, v).font = _f()
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).alignment = Alignment(vertical="top")
        ws.row_dimensions[i].height = max(15, 13 * (len(v) // 120 + 1))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 140


def _exec_summary(wb, dr_total, pr_first, findings, actions, meta):
    ws = wb["Executive Summary"]
    _title(ws, f"ENPPI Workload Forecast - Discipline Comments Analysis ({meta['cycle']})",
           f"Planning Input vs Discipline Comment vs After Implementation | window {meta['window']} | generated {meta['generated']}")
    D = "'Department Ranking'!"
    t = dr_total
    kpis = [("Planning Input MH", f"={D}D{t}", MH), ("Discipline-adjusted MH", f"={D}E{t}", MH),
            ("Net deviation requested", f"={D}I{t}", MH), ("Net deviation %", f"={D}J{t}", PCT),
            ("Gross increase requested", f"={D}G{t}", MH), ("Gross reduction requested", f"={D}H{t}", MH),
            ("After Implementation MH", f"={D}F{t}", MH), ("Implementation rate of requests", f"={D}V{t}", PCT),
            ("Concern items", f"={D}M{t}", "#,##0"), ("New project rows", f"={D}P{t}", "#,##0"),
            ("Planner change without comment MH", f"={D}W{t}", MH), ("Blank/zero not applied MH", f"={D}S{t}", MH)]
    ws.cell(4, 1, "Headline figures (24-month window)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for k, (lab, f, fmt) in enumerate(kpis):
        r, c = 5 + (k // 4) * 2, 1 + (k % 4) * 3
        ws.cell(r, c, lab).font = _f(color=MUTED)
        v = ws.cell(r + 1, c, f)
        v.font = Font(name=FONT, bold=True, size=14, color=NAVY)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 2)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 2)
        ws.row_dimensions[r + 1].height = 22
    row = 12
    ws.cell(row, 1, "Key findings").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for k, txt in enumerate(findings, start=1):
        c = ws.cell(row + k, 1, f"{k}. {txt}")
        c.font, c.alignment = _f(size=10), Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + k, start_column=1, end_row=row + k, end_column=12)
        ws.row_dimensions[row + k].height = 14 * (len(txt) // 165 + 1) + 4
    row += len(findings) + 2
    ws.cell(row, 1, "Top 10 projects by severity").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    P = "'Project Ranking'!"
    heads = ["Rank", "Project", "Name", "Severity", "Depts", "Planning MH", "Discipline MH", "Net", "Net %", "After Impl. MH", "Impl. rate"]
    _header(ws, row + 1, 1, heads)
    src = ["A", "B", "D", "E", "G", "I", "J", "N", "O", "K", "V"]
    for k in range(10):
        r = row + 2 + k
        for j, col in enumerate(src, start=1):
            c = ws.cell(r, j, f"={P}{col}{pr_first + k}")
            c.font, c.border = _f(), BOX
            c.number_format = PCT if col in ("O", "V") else (MH if col in ("I", "J", "N", "K") else "General")
    for sev, colr in SEV_FILL.items():
        ws.conditional_formatting.add(f"D{row + 2}:D{row + 11}", CellIsRule(operator="equal", formula=[f'"{sev}"'], fill=_fill(colr),
                                      font=Font(name=FONT, color="FFFFFF" if sev == "Critical" else INK)))
    row += 14
    ws.cell(row, 1, "Top 8 departments by absolute net deviation").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    heads = ["Dept", "Department", "", "Planning MH", "Discipline MH", "Net", "Dev. %", "Concern items", "Impl. rate", "Avg util. (disc.)"]
    _header(ws, row + 1, 1, heads)
    ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=3)
    for k in range(8):
        r = row + 2 + k
        for j, col in [(1, "A"), (2, "B"), (4, "D"), (5, "E"), (6, "I"), (7, "J"), (8, "M"), (9, "V"), (10, "AA")]:
            c = ws.cell(r, j, f"={D}{col}{5 + k}")
            c.font, c.border = _f(), BOX
            c.number_format = PCT if col in ("J", "V", "AA") else (MH if col in ("D", "E", "I") else "General")
        ws.cell(r, 3).border = BOX
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    row += 12
    ws.cell(row, 1, "Recommended management actions").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for k, txt in enumerate(actions, start=1):
        c = ws.cell(row + k, 1, f"{k}. {txt}")
        c.font, c.alignment = _f(size=10), Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + k, start_column=1, end_row=row + k, end_column=12)
        ws.row_dimensions[row + k].height = 14 * (len(txt) // 165 + 1) + 4
    widths = [8, 14, 40, 12, 12, 14, 14, 13, 12, 14, 12, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[L(j)].width = w


def write_workbook(path, *, rows, cells, transfers, ds, ps, mdf, cap, chk, inventory, issues, findings, actions, meta, possible=None):
    wb = Workbook()
    wb.remove(wb.active)
    for name in ["Executive Summary", "Department Ranking", "Project Ranking", "Monthly Workload Impact",
                 "Deviation Register", "Cell Changes", "New Projects Register", "Phase Transfers & Rescheduling",
                 "Implementation Tracking", "Comments Checklist", "Data Quality Log", "Method & Assumptions"]:
        wb.create_sheet(name)
    _, rr_last = _row_register(wb, rows)
    _, md_last = _monthly_data(wb, mdf)
    _, cap_last = _capacity(wb, cap)
    _, _, dr_total = _dept_ranking(wb, ds, rr_last, md_last, cap_last)
    pr_first, _ = _project_ranking(wb, ps, rr_last)
    months = sorted(cells["Month"].unique())
    dm = mdf.groupby(["DeptCode", "Month"]).sum(numeric_only=True)
    dev_lim = float(max(1.0, (dm["Discipline_Final_MH"] - dm["Planning_MH"]).abs().max()))
    _monthly_impact(wb, months, md_last, cap_last, dev_lim)
    _deviation_register(wb, rows, ps)
    _cell_changes(wb, cells, rows)
    _new_projects(wb, rows)
    _phase_and_schedule(wb, rows, transfers, issues, possible)
    _implementation(wb, rows, rr_last)
    _checklist(wb, chk)
    _dq(wb, inventory, issues)
    _method(wb, meta)
    _exec_summary(wb, dr_total, pr_first, findings, actions, meta)
    tab = {"Executive Summary": NAVY, "Row Register": "7F7F7F", "Monthly Data": "7F7F7F", "Capacity": "7F7F7F"}
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws.title in ("Monthly Workload Impact", "Department Ranking") else ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.sheet_properties.tabColor = tab.get(ws.title, "2A78D6")
        ws.sheet_view.zoomScale = 90
    wb.save(path)
