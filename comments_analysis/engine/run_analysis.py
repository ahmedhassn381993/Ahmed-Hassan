"""Run the full discipline-comments analysis for one forecast cycle.

Example (this cycle - originals not uploaded, so the corrected Copilot table fills the gaps):
    python run_analysis.py --cycle "Q3 2026" \
        --copilot-fact PowerBI_Workload_Model_3_Scenarios_2026-08-07_0106.xlsx \
        --comments COMMENTS --employees employee_count.xlsx \
        --checklist "Critical_Projects_-__Workload_Forecast_Comments_Checklist.xlsx" \
        --master PROJECTS_my_db.xlsx --tracker Enppi_Workload_Forecast_Workflow_Steps__Schedule.xlsx --out ../output

Normal use once all workbooks are available (no Copilot file needed):
    python run_analysis.py --cycle "Q4 2026" --before "To Departments" --comments COMMENTS --after after ...
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import textwrap
from pathlib import Path

import openpyxl
import pandas as pd

from analysis import (DEPARTMENTS, attach_master, checklist, company_peaks, copy_paste_check, department_summary,
                      feedback_status, load_capacity, load_master, load_tracker, monthly, possible_transfers, project_summary,
                      section_breakdown, suggest_codes)
from compare import _norm_name
from export_webapp import export_webapp
from build_dataset import build
from compare import three_way
from powerbi_export import export_powerbi
from report_excel import write_workbook


PIPING = ["2410", "2420", "2430", "2440", "2450"]


def read_checklist(path) -> list[dict]:
    """Critical projects with the planner columns (Client, Planning Engineer, Planning Supervisor)."""
    if not path or not Path(path).exists():
        return []
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out = []
    for r in range(4, ws.max_row + 1):
        code, title = ws.cell(r, 2).value, ws.cell(r, 3).value
        if code not in (None, "") and str(title).strip().lower() != "dummy":
            txt = lambda c: str(ws.cell(r, c).value or "").strip()
            out.append(dict(code=str(code).strip(), title=str(title).strip(), client=txt(4), engineer=txt(5), supervisor=txt(6)))
    return out


def file_inventory(df: pd.DataFrame, fb: pd.DataFrame) -> pd.DataFrame:
    rec = []
    f = fb.set_index("DeptCode")
    for dept, name in DEPARTMENTS.items():
        for sc in ("Planning Input", "Discipline Comment", "System Implementation"):
            g = df[(df["Scenario"] == sc) & (df["DeptCode"] == dept)]
            file = g["SourceFile"].iloc[0] if len(g) else "(missing)"
            src = g["SourceType"].iloc[0] if len(g) else "-"
            action, sent, received = "", None, None
            if not len(g):
                action = "File missing - obtain it."
            elif sc == "Discipline Comment":
                status = f.loc[dept, "Feedback_Status"]
                sent, received = f.loc[dept, "Tracker_Sent"], f.loc[dept, "Tracker_Received"]
                if "file not provided" in status:
                    action = f"Comments were received on {received:%d-%b-%y} but the -c file was not provided - upload it."
                elif status.startswith("No comments received"):
                    action = ("No comments received (tracker) - chase the department for comments or a written 'no comments'."
                              + (" Piping may return one consolidated 2400-c file; the engine splits it into 2410-2450 by its "
                                 "sub-department headers." if dept in PIPING else ""))
                elif not f.loc[dept, "Feedback_Processed"]:
                    action = "Only the original is available - confirm 'no comments' with the department."
                elif src.startswith("Copilot"):
                    action = "Upload the -c workbook to resolve blank-vs-zero cells and red/yellow markings."
            elif src.startswith("Copilot"):
                action = "Original workbook not uploaded; values taken from the Copilot table after validation."
            rec.append({"Dept": int(dept), "Department": name, "Scenario": sc, "File Used": file, "Source": src,
                        "Sent to Dept": sent if pd.notna(sent) else None,
                        "Comments Received": received if pd.notna(received) else None, "Action Needed": action})
    return pd.DataFrame(rec)


def _title_only(name) -> str:
    """Drop the leading project code from an activity title."""
    return " ".join(re.sub(r"^[\d\s,\-]+", "", str(name or "")).split())


def _short(text: str, width: int = 45) -> str:
    """Shorten a title at a word boundary for management text."""
    return textwrap.shorten(text, width=width, placeholder="...")


TITLE_STOP = {"adnoc", "project", "projects", "phase", "plant", "works", "work", "epc", "the", "for", "and", "services",
              "service", "engineering", "company", "new", "scope", "production", "development", "onshore", "offshore"}


def _title_clusters(sub: pd.DataFrame) -> list[tuple[str, float]]:
    """Group a code's activity titles that share a keyword (e.g. two wordings of the same methanol plant);
    returns (title with most MH, MH) per group, largest first."""
    groups = []
    for t, mh in sub.groupby("T")["Final_Total"].sum().sort_values(ascending=False).items():
        words = set(re.findall(r"[a-z]{3,}", t.lower())) - TITLE_STOP
        for g in groups:
            if g["words"] & words:
                g["words"] |= words
                g["MH"] += mh
                break
        else:
            groups.append({"title": t, "words": words, "MH": mh})
    return [(g["title"], g["MH"]) for g in sorted(groups, key=lambda g: -g["MH"])]


def _and(xs) -> str:
    xs = [str(x) for x in xs]
    return "" if not xs else xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


def _depts(codes) -> str:
    """Department list with the five Piping sections collapsed to 2400."""
    codes = sorted(set(codes), key=int)
    if set(PIPING) <= set(codes):
        codes = sorted(["2400 Piping (2410-2450)"] + [c for c in codes if c not in PIPING], key=lambda c: int(c[:4]))
    return _and(codes)


def build_findings(ds, ps, rows, issues, mdf, cap, note: str = "", possible=None, transfers=None,
                   peaks=None) -> tuple[list[str], list[str]]:
    f = []
    P, I = ds["Planning_MH"].sum(), ds["Implemented_MH"].sum()
    inc, red = ds["Gross_Increase"].sum(), ds["Gross_Reduction"].sum()
    net, req = inc + red, inc - red
    impl = rows.loc[rows["Concern"], "Implemented_Request_MH"].sum()
    unreq = ds["Unrequested_Gross"].sum()
    proc = ds[ds["Feedback_Processed"]]
    with_changes = ds[ds["Feedback_Status"] == "Feedback received - changes"]
    unchanged = ds[ds["Feedback_Status"] == "Feedback received - no MH change"]
    not_round = ds[ds["Feedback_Status"].str.contains("not in the tracker's review round", regex=False)]
    missing = ds[~ds["Feedback_Processed"]]
    not_recv = missing[missing["Feedback_Status"].str.startswith("No comments received")]
    not_prov = missing[missing["Feedback_Status"].str.contains("file not provided")]
    other = missing[~missing.index.isin(not_recv.index.union(not_prov.index))]
    no_file = other[other["Feedback_Status"] == "No feedback file"]
    original = other[other["Feedback_Status"] != "No feedback file"]
    has_tracker = any(c in ds and ds[c].notna().any() for c in ("Tracker_Sent", "Tracker_Received"))
    has_after = bool(rows["I_Exists"].any())
    pct = lambda x, y: x / y if y else 0.0

    # 1 headline and coverage
    t = (f"Departments asked for a net {net:+,.0f} MH ({pct(net, P):+.1%}) on a Planning Input of {P:,.0f} MH "
         f"(gross +{inc:,.0f} / {red:,.0f}). Feedback was processed for {len(proc)} of {len(ds)} departments, covering "
         f"{proc['Planning_MH'].sum():,.0f} MH ({pct(proc['Planning_MH'].sum(), P):.0%}) of Planning Input; {len(with_changes)} changed figures"
         + (f"; {_depts(unchanged['DeptCode'])} returned comments without MH changes" if len(unchanged) else "")
         + (f"; {_depts(not_round['DeptCode'])} is not in the tracker's review round and its -c file has no MH change"
            if len(not_round) else "") + ".")
    if len(missing):
        parts = []
        if len(not_recv):
            parts.append(f"the tracker shows no comments received from {_depts(not_recv['DeptCode'])}")
        if len(not_prov):
            dates = sorted({d for d in not_prov["Tracker_Received"] if pd.notna(d)})
            parts.append(f"{_depts(not_prov['DeptCode'])} sent comments ({_and(f'{d:%d-%b-%y}' for d in dates)}) that were not "
                         "provided for analysis")
        if len(original):
            parts.append(f"{_depts(original['DeptCode'])} returned only the original")
        if len(no_file):
            parts.append(f"no file was received from {_depts(no_file['DeptCode'])}"
                         + ("" if has_tracker else " (without the workflow tracker their comment status is unknown)"))
        t += (f" {len(missing)} departments ({missing['Planning_MH'].sum():,.0f} MH, {pct(missing['Planning_MH'].sum(), P):.0%}) have no "
              "comment file in the data, so their original is used with no change: " + "; ".join(parts) + ".")
    f.append(t)

    # 2 largest deviations, absolute and relative
    top = with_changes.assign(a=with_changes["Net_Deviation"].abs()).sort_values("a", ascending=False).head(4)
    rel = with_changes[with_changes["Planning_MH"] > 0].sort_values("Deviation_Pct", ascending=False)
    if len(top):
        f.append("Largest departmental deviations: " + "; ".join(
            f"{r.DeptCode} {r.Department} {r.Net_Deviation:+,.0f} MH ({r.Deviation_Pct:+.0%})" for r in top.itertuples()) + "."
                 + (f" Highest relative deviation: {rel.iloc[0].DeptCode} {rel.iloc[0].Department} {rel.iloc[0].Deviation_Pct:+.0%} "
                    f"({rel.iloc[0].Net_Deviation:+,.0f} MH)." if len(rel) else ""))
    else:
        f.append("No department changed any man-hour figure.")

    # 3 most concern items, keeping ties with third place
    pc = ds[ds["Concern_Items"] > 0].sort_values(["Concern_Items", "Concern_Projects"], ascending=False)
    if len(pc):
        pc = pc[pc["Concern_Items"] >= pc["Concern_Items"].iloc[min(2, len(pc) - 1)]]
    groups = []
    for n, g in pc.groupby("Concern_Items", sort=False):
        if len(g) == 1:
            r = g.iloc[0]
            groups.append(f"{r.DeptCode} {r.Department} ({int(n)} items, {int(r.Concern_Projects)} projects)")
        else:
            groups.append(_and(f"{r.DeptCode} {r.Department}" for r in g.itertuples())
                          + f" ({int(n)} items each; {_and(str(int(x)) for x in g['Concern_Projects'])} projects)")
    f.append("Most concern items: " + "; ".join(groups) + "." if groups else "No concern items were raised.")

    # 4 severity
    crit = ps[ps["Severity"] == "Critical"]
    tiers = []
    for n, g in ps[ps["Depts_Concerned"] > 0].sort_values("Gross_Increase", ascending=False).groupby("Depts_Concerned", sort=False):
        tiers.append((int(n), g))
    tiers = sorted(tiers, key=lambda x: -x[0])[:2]
    parts = []
    for n, g in tiers:
        lead = g.head(2)
        names = _and(f"{r.ProjectFull} ({r.Net_Change:+,.0f} MH)" for r in lead.itertuples())
        parts.append(f"{len(g)} project{'s' if len(g) > 1 else ''} raised by {n} departments"
                     + (f": {names}" if len(g) <= 2 else f", led by {names}"))
    f.append(f"{len(crit)} projects rated Critical and {int((ps['Severity'] == 'High').sum())} High."
             + (" Broadest concern: " + "; then ".join(parts) + "." if parts else ""))

    # 5 new allocations, with title variants, master flags and copy-paste caveats
    newp = rows[rows["Concern_Types"].str.contains("New Project|New Phase")]
    big_new = newp.groupby("ProjectFull").agg(MH=("Final_Total", "sum"), D=("DeptCode", "nunique")).sort_values("MH", ascending=False).head(4)
    copy = [i for i in issues if i["Area"] == "Possible copy-paste"]
    parts = []
    for code, v in big_new.iterrows():
        sub = newp[newp["ProjectFull"] == code].assign(T=lambda x: x["D_Name"].map(_title_only))
        tg = _title_clusters(sub)
        d = f"{int(v.D)} dept{'s' if v.D > 1 else ''}"
        if len(tg) == 1:
            txt = f"{code} {_short(tg[0][0])} ({v.MH:,.0f} MH, {d})"
        else:
            txt = (f"{code} ({v.MH:,.0f} MH, {d}) used for {len(tg)} different works ("
                   + "; ".join(f"{_short(t, 40)} {mh:,.0f} MH" for t, mh in tg) + ")")
        flag = str(sub["Master_Flag"].iloc[0]) if "Master_Flag" in sub else ""
        if flag.startswith("Internal / overhead"):
            txt += f"; the project master lists {code} as '{str(sub['Master_Name'].iloc[0]).strip().title()}', so a project code is still to be assigned"
        elif flag.startswith("Not in project master"):
            txt += ", not yet in the project master"
        cp = [c for c in copy if code in c["Detail"]]
        if cp:
            txt += f"; its monthly profile repeats another project's in {cp[0]['File']} (possible copy-paste - confirm)"
        parts.append(txt)
    n_phase = int(newp["Primary_Type"].eq("New Phase Allocation").sum())
    if not len(newp):
        f.append("No new allocations were requested.")
    else:
        f.append(f"{len(newp)} new allocations ({len(newp) - n_phase} new projects, {n_phase} new phase allocations; "
             f"{newp['Final_Total'].sum():,.0f} MH) were requested; the largest are " + "; ".join(parts) + "."
             + (f" Only {int(newp['I_Exists'].sum())} of the {len(newp)} appear in the after-implementation data." if has_after else ""))

    # 6 implementation
    gap_dep = ds[~ds["Feedback_Processed"]]
    gap = gap_dep["Unrequested_Gross"].sum()
    gap_prov = not_prov["Unrequested_Gross"].sum()
    silent = missing[~missing.index.isin(not_prov.index)]
    gap_silent = silent["Unrequested_Gross"].sum()
    big_silent = silent.sort_values("Unrequested_Gross", ascending=False).iloc[0] if len(silent) else None
    unreq_net = ds["Unrequested_Net"].sum()
    share = "essentially all" if abs(I - P) > 0 and abs(unreq_net - (I - P)) <= 0.05 * abs(I - P) else "most"
    if not has_after:
        t = ("Implementation: no after-implementation extract was provided, so implementation is not measured yet"
             + (f" ({req:,.0f} MH of change was requested)." if req >= 1 else "."))
    else:
        t = (f"Implementation: of {req:,.0f} MH of requested change, about {impl:,.0f} MH ({pct(impl, req):.0%}) appears in the "
             "after-implementation extract." if req >= 1 else "Implementation: no man-hour change was requested.")
        if unreq >= 1:
            t += (f" Separately, the extract differs from Planning Input by {unreq:,.0f} MH (gross) in cells where no department "
                  "asked for a change"
                  + (f", so {share} of the {I - P:+,.0f} MH movement is not explained by comments" if abs(I - P) >= 1 else "")
                  + " and should be reconciled before the extract is used as 'final'."
                  + (f" Of that gross movement, {gap_prov:,.0f} MH is in {_depts(not_prov['DeptCode'])}, whose comments were sent but "
                     "not provided (it may be implemented comments)," if gap_prov >= 1 else " Of that gross movement,")
                  + (f" {gap_silent:,.0f} MH is in departments that sent no comments ({_depts(silent['DeptCode'])}; largest "
                     f"{big_silent.DeptCode}, {big_silent.Unrequested_Gross:,.0f} MH), so it is planner change without any comment, and "
                     if gap_silent >= 1 and big_silent is not None else "")
                  + f" {unreq - gap:,.0f} MH is in departments whose comments were processed.")
        else:
            t += " Outside the requested cells the extract equals Planning Input."
    f.append(t + (f" {note}" if note else ""))

    # 7 schedule extensions
    ext = rows[rows["Concern_Types"].str.contains("Schedule Extension")]
    ext_top = ext.groupby("ProjectFull").agg(MH=("Extension_MH", "sum"), D=("DeptCode", "nunique")).sort_values("MH", ascending=False).head(4)
    if len(ext):
        f.append(f"{len(ext)} rows extend the schedule beyond the baseline finish ({ext['Extension_MH'].sum():,.0f} MH falls after "
                 "the planned end). Largest: " + "; ".join(f"{k} ({v.MH:,.0f} MH, {int(v.D)} dept{'s' if v.D > 1 else ''})"
                                                         for k, v in ext_top.iterrows()) + ".")

    # 8 Loan To
    loan = rows[rows["Concern"] & (rows["Phase"] == "Loan To")]
    if len(loan):
        f.append(f"{len(loan)} Loan To rows were changed by {loan['DeptCode'].nunique()} departments (gross "
                 f"{(loan['Gross_Increase'] - loan['Gross_Reduction']).sum():,.0f} MH); Loan To is Corporate Planning controlled, "
                 "so these are requests, not edits.")

    # 9 transfers: same code across phases, and reductions paired with new codes
    tr = transfers if transfers is not None else pd.DataFrame()
    n_pt = len(possible) if possible is not None else 0
    if n_pt or len(tr):
        off = (possible["Possible Offset MH"].sum() if n_pt else 0) + (tr["Transfer_MH"].sum() if len(tr) else 0)
        items = []
        if n_pt:
            for _, t in possible.sort_values("Possible Offset MH", ascending=False).iterrows():
                items.append(f"{t['Dept']}: {t['Reduced Project']} {t['Reduced Phase']} {t['Reduction MH']:,.0f} MH against new "
                             f"{t['Added Project']} {t['Added Phase']} +{t['Added MH']:,.0f} MH")
        for t in tr.itertuples():
            sub = rows[(rows["DeptCode"] == t.DeptCode) & (rows["ProjectBase"] == t.ProjectBase) & (rows["Net_Change"].abs() >= 1)]
            if not ((sub["Net_Change"] < 0).any() and (sub["Net_Change"] > 0).any()):
                continue
            dn = sub[sub["Net_Change"] < 0].iloc[0]
            up = sub[sub["Net_Change"] > 0].iloc[0]
            items.append(f"{t.DeptCode}: {dn.ProjectFull} {dn.Phase} {dn.Net_Change:,.0f} MH moved to {up.ProjectFull} {up.Phase} "
                         f"+{up.Net_Change:,.0f} MH")
        f.append(f"{n_pt + len(tr)} reductions look like transfers to another code or phase in the same department, so up to "
                 f"{off:,.0f} MH is counted twice in the gross figures: " + "; ".join(items) + ".")

    # 10 code corrections
    cc = rows[rows["CodeCorrection"] != ""]
    if len(cc):
        notes = []
        for r in cc.itertuples():
            old, new = r.CodeCorrection.split(" (")[0].split(" -> ")
            if "typed in column A" not in r.CodeCorrection:
                # as new work the new code adds its full value and the baseline row stays: net rises by the baseline
                notes.append(f"{r.DeptCode} re-coded {old} as {new} ({r.Final_Total:,.0f} MH); it is paired with its baseline - "
                             f"counted as new work with the baseline retained, the company net would be "
                             f"{net + r.P_Total:+,.0f} MH ({r.P_Total:+,.0f})")
            elif r.P_Exists:
                notes.append(f"{r.DeptCode} typed {new} in column A on the {old} row ({r.Final_Total:,.0f} MH); the activity code is "
                             f"kept - if {new} is confirmed, that MH belongs to {new} instead (as new work with the "
                             f"{r.P_Total:,.0f} MH baseline retained, the net would rise by a further {r.P_Total:,.0f} MH)")
            else:
                notes.append(f"{r.DeptCode} typed {new} in column A on a new {old} row ({r.Final_Total:,.0f} MH), "
                             "a new allocation under either code")
        f.append(f"{len(cc)} rows carry a code change by the department: " + "; ".join(notes) + ". Confirm each code in Primavera.")

    # 11 phase moves
    moves = rows[rows["RequestedPhase"] != ""]
    if len(moves):
        f.append(f"{len(moves)} projects ({moves['P_Total'].sum():,.0f} MH) appear under a different phase in the feedback "
                 f"(departments {_depts(moves['DeptCode'])}); these were paired, not added, "
                 "which prevents double counting. Category changes need Corporate Planning approval.")

    # 12 Copilot model
    dup = sum(i["Impact_MH"] or 0 for i in issues if i["Area"] == "Double counting")
    hdr = sum(i["Impact_MH"] or 0 for i in issues if i["Issue"].startswith("Department / sub-department total rows"))
    nonmonth = sorted({i["File"] for i in issues if i["Issue"].startswith("Non-month column")})
    direct = sorted(rows.loc[rows["D_SourceType"].fillna("").eq("Source workbook"), "D_File"].dropna().unique())
    summary = next((i for i in issues if i["Area"] == "Copilot check" and i["Severity"] == "Info"), None)
    differ = [i for i in issues if i["Area"] == "Copilot check" and i["Severity"] != "Info"]
    n_checked = int(summary["Detail"].split(" ")[0]) if summary else 0
    if n_checked:
        k = sum(int(i["Detail"].split(" of ")[0].replace(",", "")) for i in differ)
        check = (f"Compared project by month with the {n_checked} department files read directly, it "
                 + (f"differs in {_and(sorted({i['File'] for i in differ}, key=int))} ({k:,} cell{'s' if k != 1 else ''}; "
                    "see Data Quality Log)" if differ else "agrees")
                 + "; the files are used there, and the corrected table only where no file was provided.")
    else:
        check = "Its values were reused after correction."
    if hdr or dup or nonmonth:
        f.append("The Copilot 3-scenario model is not reliable as built: it counted department total rows as projects "
                 f"({hdr:,.0f} MH across scenarios)"
                 + (f", a non-month column in {', '.join(nonmonth)} as a month" if nonmonth else "")
                 + f", and after-implementation blocks that repeat other departments ({dup:,.0f} MH). " + check)

    # 13 capacity: averages, company peak, months over capacity
    util = ds.dropna(subset=["Firm_Utilisation_Final"]).sort_values("Firm_Utilisation_Final", ascending=False).head(5)
    dm = mdf.groupby(["DeptCode", "Month"])["Discipline_Final_MH"].sum().reset_index().merge(cap[["DeptCode", "Capacity_MH_per_Month"]], on="DeptCode")
    over = dm[dm["Discipline_Final_MH"] > dm["Capacity_MH_per_Month"]]
    months = mdf["Month"].nunique()
    t = ("Highest firm utilisation (Running + Internal + Loan To + Under Signature, discipline-adjusted): " + "; ".join(
        f"{r.DeptCode} {r.Firm_Utilisation_Final:.0%}" for r in util.itertuples()) + ".")
    if peaks:
        pk, pi = peaks["Discipline_Final_MH"], peaks["Implemented_MH"]
        t += (f" Company workload peaks in {pk['Month']:%b-%y} at {pk['MH']:,.0f} MH ({pk['Utilisation']:.0%} of capacity) "
              f"discipline-adjusted, and in {pi['Month']:%b-%y} at {pi['MH']:,.0f} MH ({pi['Utilisation']:.0%}) after implementation.")
    fo = ds[ds["Firm_Overload_Months_Final"] > 0].sort_values("Firm_Overload_Months_Final", ascending=False)
    if len(fo):
        grp = [f"{_and(g['DeptCode'])} ({int(n)})" for n, g in fo.head(6).groupby("Firm_Overload_Months_Final", sort=False)]
        t += f" Firm workload alone exceeds capacity in the most months (of {months}) in {', '.join(grp)}."
    t += (f" Including Proposal and Potential, {over['DeptCode'].nunique()} departments exceed capacity in at least one month "
          "- see 'Monthly Workload Impact'.")
    f.append(t)

    # ---------------------------------------------------------------- actions
    crit_codes = ps[ps["Severity"] == "Critical"].sort_values("Gross_Increase", ascending=False)
    new_crit = [c for c in crit_codes["ProjectFull"] if c in set(big_new.index)][:2]
    widest = ps[ps["Depts_Concerned"] > 0].sort_values(["Depts_Concerned", "Gross_Increase"], ascending=False)["ProjectFull"].head(2).tolist()
    first = "Critical" if len(crit_codes) else "High" if (ps["Severity"] == "High").any() else ""
    placeholder = (crit_codes.loc[crit_codes["Master_Flag"].fillna("").str.startswith("Internal / overhead"), "ProjectFull"].tolist()
                   if "Master_Flag" in crit_codes else [])
    a = []
    plural = lambda xs, one, many: f"{one if len(xs) == 1 else many} {_and(xs)}"
    if rows["Concern"].any():
        a.append((f"Corporate Planning to rule on the {first} projects first (see Project Ranking)" if first else
                  "Corporate Planning to rule on each request in Project Ranking order")
                 + ": accept, reduce or return each change with a reason"
                 + ((", starting with " + (plural(new_crit, "the new project", "the new projects") + " and " if new_crit else "")
                     + plural(widest, "the most widely raised project", "the most widely raised projects")) if widest else "")
                 + (f"; assign a proper project code to {_and(placeholder)} (an internal account in the project master)" if placeholder else "")
                 + ".")
    if has_after and rows["Implementation_Status"].str.startswith("Not implemented").any():
        a.append("Planners to close every 'Not implemented' item in Implementation Tracking: implement it or mark it 'not applicable' "
                 "with a reason, so the checklist no longer depends on email follow-up.")
    if has_after and unreq >= 1:
        a.append("Reconcile the after-implementation extract: explain the MH that changed without a discipline comment before the "
                 "forecast is locked" + (", starting with the departments whose comments were not processed." if gap >= 1 else "."))
    if len(missing):
        parts = []
        if len(not_prov):
            parts.append(f"obtain the comment files that {_depts(not_prov['DeptCode'])} returned "
                         f"({_and(sorted({f'{d:%d-%b-%y}' for d in not_prov['Tracker_Received'] if pd.notna(d)}))})")
        if len(not_recv):
            parts.append(f"chase {_depts(not_recv['DeptCode'])} for comments or a written 'no comments'")
        if len(original):
            parts.append(f"confirm 'no comments' from {_depts(original['DeptCode'])} (they returned the original)")
        if len(no_file):
            parts.append(f"obtain the feedback, or a written 'no comments', from {_depts(no_file['DeptCode'])}")
        txt = "; ".join(parts)
        a.append(txt[0].upper() + txt[1:] + f", so the analysis covers all {len(ds)} departments.")
    cat = []
    for dept, g in moves.groupby("DeptCode"):
        cat.append(f"{dept} ({len(g)} rows)")
    for t in tr.itertuples():
        cat.append(f"{t.DeptCode} ({t.ProjectBase}: {t.phases_down} -> {t.phases_up})")
    txt = ""
    if cat:
        txt = "the category changes requested by " + _and(cat)
    if n_pt:
        txt += ("; also " if txt else "") + "the possible transfers in " + _and(
            f"{t['Dept']} ({t['Reduced Project']} {t['Reduced Phase']} -> {t['Added Project']} {t['Added Phase']})"
            for _, t in possible.iterrows())
    if len(loan):
        txt += (" and " if txt else "") + (f"the {len(loan)} Loan To changes, which are Corporate Planning controlled" if len(loan) > 1
                                           else "the Loan To change, which is Corporate Planning controlled")
    if txt:
        a.append("Confirm " + txt + ".")
    if len(cc):
        a.append("Confirm the department code changes in Primavera before the next cycle: "
                 + _and(f"{r.DeptCode} {r.CodeCorrection.split(' (')[0]}" for r in cc.itertuples()) + ".")
    dept_areas = {"Source workbook", "Duplicate project row", "Project code", "Phase transfer", "Possible copy-paste",
                  "Possible transfer"}
    dept_issues = [i for i in issues if i["Area"] in dept_areas]
    fixable = [i for i in dept_issues if re.search(r"code conflict|Uncoded|code corrected|Phase heading|Row-total|different phase",
                                                    i["Issue"], re.I)]
    copilot = [i for i in issues if i["Area"] in ("Copilot model", "Double counting")]
    a.append("Fix the template: replace =LEFT(B,7) with the full code, lock phase headings, forbid row totals inside month columns and "
             f"give departments a 'Reason' column. These remove the causes of {len(fixable)} of the {len(dept_issues)} issues logged "
             f"from department files (code conflicts, uncoded rows, renumbered or missing phase headings, a typed row total); "
             f"{len(copilot)} of the {len(issues)} logged issues came from how the Copilot model was built.")
    hot = ds[ds["Firm_Overload_Months_Final"] >= max(1, months // 3)].sort_values("Firm_Overload_Months_Final", ascending=False)
    a.append("Review hiring or redistribution for departments whose firm workload stays above capacity for several months"
             + (": " + _and(f"{r.DeptCode} ({int(r.Firm_Overload_Months_Final)}" + (f" of {months} months)" if k == 0 else ")")
                            for k, r in enumerate(hot.itertuples())) if len(hot) else "") + ".")
    tidy = lambda xs: [re.sub(r" {2,}", " ", x) for x in xs]
    return tidy(f), tidy(a)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cycle", default="Forecast cycle")
    ap.add_argument("--copilot-fact")
    ap.add_argument("--before")
    ap.add_argument("--comments")
    ap.add_argument("--after")
    ap.add_argument("--employees", required=True)
    ap.add_argument("--checklist")
    ap.add_argument("--master", help="project master (PROJECTS_my_db.xlsx) for names, clients and code checks")
    ap.add_argument("--tracker", help="workflow tracker with the 'Disciplines Review' sent/received dates")
    ap.add_argument("--out", required=True)
    ap.add_argument("--note", default="", help="context sentence added to the implementation finding")
    ap.add_argument("--webapp-export", help="also write import files for the Workload Forecast web application here")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    folders = {"Planning Input": a.before, "Discipline Comment": a.comments, "System Implementation": a.after}
    df, issues = build(a.copilot_fact, folders)
    cells, rows, transfers = three_way(df, issues)
    rows["Implemented_Request_MH"] = rows["Impl_Ratio"].fillna(0) * (rows["Gross_Increase"] - rows["Gross_Reduction"]) * rows["Concern"]
    master = load_master(a.master)
    rows = attach_master(rows, "ProjectFull", master)
    issues += copy_paste_check(df)
    for s in suggest_codes(rows).itertuples():
        issues.append(dict(Severity="Medium", Area="Project code", Scenario="Discipline Comment", File=s.DeptCode,
                           Issue="Uncoded feedback row - closest known project (verify)",
                           Detail=f"'{s.Activity}' -> {s.Suggested_Code or 'no close match'} ({s.Match:.0%})", Impact_MH=s.Final_MH))
    possible = possible_transfers(rows)
    for _, t in possible.iterrows():
        issues.append(dict(Severity="High" if t["Possible Offset MH"] >= 5000 else "Medium", Area="Possible transfer",
                           Scenario="Discipline Comment", File=str(t["Dept"]),
                           Issue="Reduction paired with a new allocation of the same family or title (verify)",
                           Detail=f"{t['Reduced Project']} {t['Reduced Phase']} {t['Reduction MH']:,.0f} MH vs {t['Added Project']} "
                                  f"{t['Added Phase']} +{t['Added MH']:,.0f} MH ({t['Match Basis']})",
                           Impact_MH=float(t["Possible Offset MH"])))
    if len(master):
        flagged = rows[rows["Concern"] & rows["Master_Flag"].isin(["Internal / overhead account used for project work",
                                                                   "Not in project master (new or proposal code)"])]
        for (code, flag), g in flagged.groupby(["ProjectFull", "Master_Flag"]):
            internal = flag.startswith("Internal")
            issues.append(dict(Severity="Medium" if internal else "Low", Area="Project master", Scenario="Discipline Comment",
                               File=", ".join(sorted(g["DeptCode"].unique(), key=int)),
                               Issue=("Internal / BD account used for project work - proposal placeholder, code to be assigned"
                                      if internal else "Code not in the project master (new or proposal code - verify)"),
                               Detail=f"{code}: master '{g['Master_Name'].iloc[0]}' / feedback '{g['D_Name'].fillna(g['P_Name']).iloc[0]}'"
                               if internal else f"{code}: '{g['D_Name'].fillna(g['P_Name']).iloc[0]}'",
                               Impact_MH=float(g["Final_Total"].sum())))
    cap = load_capacity(a.employees)
    fb = feedback_status(rows, df, load_tracker(a.tracker))
    ds = department_summary(rows, cells, cap, fb)
    ps = project_summary(rows, cells, master if len(master) else None)
    mdf = monthly(cells, rows)
    peaks = company_peaks(mdf, cap)
    chk = checklist(rows, read_checklist(a.checklist), fb)
    inv = file_inventory(df, fb)
    sections = section_breakdown(rows)
    months = sorted(cells["Month"].unique())
    window = f"{months[0]:%b-%y} to {months[-1]:%b-%y} ({len(months)} months)"
    direct = {sc: sorted(df.loc[(df["Scenario"] == sc) & (df["SourceType"] == "Source workbook"), "DeptCode"].unique(), key=int)
              for sc in ("Planning Input", "Discipline Comment", "System Implementation")}
    copilot = {sc: sorted(set(df.loc[df["Scenario"] == sc, "DeptCode"]) - set(direct[sc]), key=int) for sc in direct}
    src = []
    for sc in direct:
        bits = []
        if direct[sc]:
            bits.append(f"{len(direct[sc])} department files read directly ({_depts(direct[sc])})")
        if copilot[sc]:
            bits.append(f"{len(copilot[sc])} departments from the corrected Copilot table")
        src.append(f"{sc}: " + "; ".join(bits) if bits else f"{sc}: none")
    if a.copilot_fact:
        src.append(f"Copilot table: {Path(a.copilot_fact).name} (compared project by month with every file read directly; "
                   "differences are in the Data Quality Log)")
    for label, path in (("Headcount", a.employees), ("Project master", a.master), ("Workflow tracker", a.tracker),
                        ("Critical projects checklist", a.checklist)):
        if path:
            src.append(f"{label}: {Path(path).name}")
    missing = fb[~fb["Feedback_Processed"]]
    lim = []
    if copilot["Planning Input"]:
        lim.append(f"Original 'To Departments' workbooks were not uploaded for {len(copilot['Planning Input'])} departments; their "
                   "Planning Input comes from the corrected Copilot table")
    if copilot["System Implementation"]:
        lim.append("After Implementation comes from the Copilot table for all departments (no 'after' workbooks uploaded)"
                   if not direct["System Implementation"] else
                   f"After Implementation comes from the Copilot table for {len(copilot['System Implementation'])} departments")
    if len(missing):
        lim.append(f"No comment file was available for {_depts(missing['DeptCode'])}; their original is used with no change")
    lim.append("Headcount date of 'Table for Emails' unconfirmed")
    processed = set(fb.loc[fb["Feedback_Processed"].astype(bool), "DeptCode"].astype(str))
    cop_fb = [d for d in copilot["Discipline Comment"] if str(d) in processed]
    cop_orig = [d for d in copilot["Discipline Comment"] if str(d) not in processed]
    meta = dict(cycle=a.cycle, window=window, generated=dt.date.today().strftime("%d-%b-%Y"), sources=" | ".join(src),
                limitations=". ".join(lim) + ".", n_direct_feedback=len(direct["Discipline Comment"]),
                n_copilot_feedback=len(copilot["Discipline Comment"]),
                copilot_feedback_text=", ".join(x for x in (
                    f"{len(cop_fb)} departments' feedback ({_depts(cop_fb)})" if cop_fb else "",
                    f"{len(cop_orig)} originals" if cop_orig else "") if x) + " this cycle" if copilot["Discipline Comment"] else "none this cycle")
    paired = rows[(rows["CodeCorrection"] != "") & ~rows["CodeCorrection"].str.contains("typed in column A")]
    meta["bridge_memo"] = [
        ("possible transfers inside the gross figures (counted in both an increase and a reduction)",
         float((possible["Possible Offset MH"].sum() if len(possible) else 0) + (transfers["Transfer_MH"].sum() if len(transfers) else 0))),
        ("code corrections paired with their baseline (would add this if counted as new work)", float(paired["P_Total"].sum()))]
    findings, actions = build_findings(ds, ps, rows, issues, mdf, cap, a.note, possible, transfers, peaks)
    tag = a.cycle.replace(" ", "_")
    write_workbook(out / f"ENPPI_Workload_Comments_Analysis_{tag}.xlsx", rows=rows, cells=cells, transfers=transfers,
                   ds=ds, ps=ps, mdf=mdf, cap=cap, chk=chk, inventory=inv, issues=issues, findings=findings,
                   actions=actions, meta=meta, possible=possible, sections=sections)
    # Power BI carries every double-count candidate: code/title pairs plus same-code phase transfers (section B)
    phase_pairs = []
    for t in transfers.itertuples():
        sub = rows[(rows["DeptCode"] == t.DeptCode) & (rows["ProjectBase"] == t.ProjectBase) & (rows["Net_Change"].abs() >= 1)]
        dn, up = sub[sub["Net_Change"] < 0], sub[sub["Net_Change"] > 0]
        if len(dn) and len(up):
            dn, up = dn.iloc[0], up.iloc[0]
            phase_pairs.append({"Dept": int(t.DeptCode), "Reduced Project": dn.ProjectFull, "Reduced Phase": dn.Phase,
                                "Reduced Title": dn.D_Name, "Reduction MH": dn.Net_Change, "Added Project": up.ProjectFull,
                                "Added Phase": up.Phase, "Added Title": up.D_Name, "Added MH": up.Net_Change,
                                "Match Basis": f"same code {t.ProjectBase}, phase transfer", "Possible Offset MH": t.Transfer_MH,
                                "Net of the Pair MH": dn.Net_Change + up.Net_Change})
    all_pairs = pd.concat([possible, pd.DataFrame(phase_pairs)], ignore_index=True) if phase_pairs else possible
    export_powerbi(out / "powerbi", df=df, cells=cells, rows=rows, ps=ps, cap=cap, issues=issues, cycle=a.cycle,
                   possible=all_pairs, fb=fb, inventory=inv)
    pd.DataFrame(issues).to_csv(out / "issues_log.csv", index=False)
    if a.webapp_export:
        print("web application import files:", export_webapp(Path(a.webapp_export), rows=rows, cells=cells, cap=cap, fb=fb))
    print("\n".join(findings))
    print("\nACTIONS\n" + "\n".join(actions))
    print(f"\nWritten to {out}")


if __name__ == "__main__":
    main()
