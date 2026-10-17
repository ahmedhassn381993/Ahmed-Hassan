"""Parse ENPPI department workload workbooks (To-Departments, -c feedback, after-implementation).

Each sheet layout (produced by the department-split macro):
    Row 1      : 'Project' | 'Activity Name' | month headers (dates) ...
    Dept row   : <dept code> | <DEPT NAME>          (totals, ignored)
    Sub-dept   : <sub code or blank> | <SUB NAME>   (totals, ignored)
    Phase row  : blank | '1.  Running' / '2. Loan To' / '3. Under Signature' / '4. Proposed' / '5. Potential Opportunity'
    Project row: =LEFT(B,7) or typed code | '<code> <client>, <title>' | monthly man-hours

Output is one record per (row, month) with the raw cell state kept, so blank cells are never
turned into zeros and the department's red/yellow/orange marks are preserved for the audit trail.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

PHASES = {
    "1": "Running",
    "2": "Loan To",
    "3": "Under Signature",
    "4": "Proposal",
    "5": "Potential Opportunity",
}
PHASE_RE = re.compile(
    r"^\s*(\d)\s*[.\-)]\s*(running|loan\s*to|under\s*signature|proposed|proposal|potential\s*opportunity|oppo)",
    re.I,
)
# the heading's words decide the phase: departments renumber headings (2300-c has "3. Proposed")
PHASE_WORDS = [("running", "Running"), ("loan", "Loan To"), ("under", "Under Signature"),
               ("propos", "Proposal"), ("potential", "Potential Opportunity"), ("oppo", "Potential Opportunity")]
# rows departments add under the projects (capacity checks, headcount); the rest of the sheet is ignored
FOOTER_RE = re.compile(
    r"(?i)^\s*(total\s*mhrs?|available\s+man|remaining|def+iciency|required\s+man\s*power|average\s+required|"
    r"no\.?\s*of\s+employees|difference|division\s*/\s*depart|qty\s+of)")
INTERNAL_CODE = "1100100"
MONTH_TEXT_RE = re.compile(r"^([A-Za-z]{3})[-\s/]?(\d{2,4})$")


@dataclass
class ProjectCode:
    full: str            # full code incl. suffix, e.g. 5020250-7, 40061420, 1251100-C1-2
    base: str            # parent code, e.g. 5020250
    suffix: str          # '' or '7' / '1' / 'C1-2'
    source: str          # 'activity' | 'column_a' | 'none'
    note: str = ""


def _digits(s: str) -> str:
    return re.sub(r"\D", "", s)


def parse_project_code(col_a, activity: str) -> ProjectCode:
    """Return the full project code without truncating 8-digit or suffixed codes.

    The template's =LEFT(B,7) formula in column A truncates 40061420 -> 4006142 and
    5020250-7 -> 5020250, so the activity text is the primary source and column A the fallback.
    """
    text = (activity or "").strip()
    for source, candidate in (("activity", text), ("column_a", str(col_a or "").strip())):
        if not candidate or candidate.startswith("="):
            continue
        # 'XXXX-YYY' / 'XXXXX-YYY' (ENPPI short form) -> merged 7/8-digit code
        m = re.match(r"^(\d{4,5})\s*-\s*(\d{3})(?!\d)", candidate)
        if m:
            code = m.group(1) + m.group(2)
            return ProjectCode(code, code, "", source, f"short form '{m.group(0)}' merged")
        # 7-8 digit code with optional suffix: 5020250-7, 3835210-1, 1251100- C1-2
        m = re.match(r"^(\d{7,8})(?:\s*-\s*([A-Za-z]?\d{1,2}(?:-\d{1,2})?)(?![\d]))?", candidate)
        if m:
            base, suffix = m.group(1), (m.group(2) or "")
            full = f"{base}-{suffix}" if suffix else base
            return ProjectCode(full, base, suffix, source)
        if source == "column_a":
            d = _digits(candidate)
            if 7 <= len(d) <= 8:
                return ProjectCode(d, d, "", source, f"digits taken from '{candidate}'")
    return ProjectCode("", "", "", "none", "no project code found")


def _month_from_header(v):
    if isinstance(v, dt.datetime):
        return dt.date(v.year, v.month, 1)
    if isinstance(v, dt.date):
        return dt.date(v.year, v.month, 1)
    if isinstance(v, (int, float)) and 40000 < v < 60000:  # Excel serial
        d = dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
        return dt.date(d.year, d.month, 1)
    if isinstance(v, str):
        m = MONTH_TEXT_RE.match(v.strip())
        if m:
            try:
                d = dt.datetime.strptime(f"{m.group(1)}-{m.group(2)[-2:]}", "%b-%y")
                return dt.date(d.year, d.month, 1)
            except ValueError:
                return None
    return None


def _rgb(color) -> str:
    try:
        if color is not None and color.type == "rgb" and color.rgb:
            return str(color.rgb)[-6:].upper()
    except Exception:  # openpyxl raises on malformed colour objects
        pass
    return ""


def _is_red(hex6: str) -> bool:
    if len(hex6) != 6:
        return False
    r, g, b = int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16)
    return r >= 0xB0 and g <= 0x50 and b <= 0x50


def _fill(cell) -> str:
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return ""
    return _rgb(f.fgColor)


@dataclass
class CellRecord:
    file: str
    sheet: str
    row: int
    dept_file: str
    section_code: str
    phase: str
    category: str
    project_full: str
    project_base: str
    project_suffix: str
    code_source: str
    code_note: str
    col_a: str
    activity: str
    month: dt.date
    raw: object
    value: float | None        # None = blank (no revision), never coerced to 0
    value_note: str
    font_red: bool
    cell_fill: str
    row_fill: str
    has_comment: bool
    comment: str = ""


@dataclass
class ParsedWorkbook:
    path: Path
    dept_file: str
    records: list[CellRecord] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    months: list[dt.date] = field(default_factory=list)


def dept_from_filename(path: Path) -> str:
    m = re.match(r"^(\d{3,4})", path.stem.split("_")[-1]) or re.search(r"(\d{3,4})(?:-c)?$", path.stem)
    return m.group(1) if m else path.stem


def parse_workbook(path: str | Path, dept_code: str | None = None) -> ParsedWorkbook:
    path = Path(path)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    out = ParsedWorkbook(path, dept_code or dept_from_filename(path))
    for ws in wb_f.worksheets:
        wsv = wb_v[ws.title]
        month_cols = {}
        for c in range(3, ws.max_column + 1):
            mth = _month_from_header(ws.cell(1, c).value)
            if mth:
                month_cols[c] = mth
        if not month_cols:
            continue
        out.months = sorted(set(out.months) | set(month_cols.values()))
        if ws.title.strip() != out.dept_file:
            out.issues.append(dict(file=path.name, sheet=ws.title, row=None, severity="Info",
                                   issue="Sheet name differs from file department code",
                                   detail=f"sheet '{ws.title}' vs file '{out.dept_file}'"))
        section, phase = out.dept_file, ""
        for r in range(2, ws.max_row + 1):
            a, b = ws.cell(r, 1), ws.cell(r, 2)
            a_val = wsv.cell(r, 1).value if isinstance(a.value, str) and a.value.startswith("=") else a.value
            text = str(b.value).strip() if b.value is not None else ""
            if text == "None":   # literal text left by copy-paste
                text = ""
            if not text and (a.value is None or str(a.value).startswith("=")):
                if all(ws.cell(r, c).value in (None, "") for c in month_cols):
                    continue  # empty template row
            pm = PHASE_RE.match(text)
            if pm and not re.match(r"^\s*\d{4,}", text):
                label = pm.group(2).lower()
                new_phase = next(v for k, v in PHASE_WORDS if label.startswith(k))
                if new_phase != PHASES[pm.group(1)]:
                    out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Warning",
                                           issue="Phase heading renumbered by department",
                                           detail=f"'{text}' read as {new_phase} (standard number {next(k for k, v in PHASES.items() if v == new_phase)})"))
                phase = new_phase
                continue
            if FOOTER_RE.match(text) and not re.match(r"^\s*\d{4,}", text):
                out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Info",
                                       issue="Department summary block below the projects ignored",
                                       detail=f"from row {r}: '{text[:60]}'"))
                break
            typed_a = a.value if not str(a.value or "").startswith("=") else None
            code = parse_project_code(typed_a, text)
            if typed_a is not None and code.source == "activity":
                a_code = parse_project_code(typed_a, "")
                if a_code.full and a_code.full != code.full and a_code.base != code.base:
                    code.note = f"column A '{typed_a}' differs from activity code {code.full}"
                    out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Warning",
                                           issue="Project code conflict between column A and activity name",
                                           detail=f"A='{typed_a}' activity='{text[:60]}' -> used {code.full}"))
            is_header = (not code.full) and (
                (isinstance(a.value, (int, float)) and a.value < 10000)
                or (a.value is None and text.upper() == text and not re.search(r"\d{5,}", text))
            )
            if is_header:
                if isinstance(a.value, (int, float)):
                    section = str(int(a.value))
                continue
            if not phase:
                out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Warning",
                                       issue="Row before any phase heading", detail=text[:80]))
                continue
            if not code.full:
                out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Warning",
                                       issue="Project code not identified", detail=f"A='{a.value}' B='{text[:60]}'"))
            category = "Internal / Unforeseen" if (phase == "Running" and code.base == INTERNAL_CODE) else phase
            row_fill = _fill(a) or _fill(b)
            for c, mth in month_cols.items():
                cf, cv = ws.cell(r, c), wsv.cell(r, c)
                raw = cf.value
                val, note = None, ""
                if isinstance(raw, str) and raw.startswith("="):
                    val = cv.value if isinstance(cv.value, (int, float)) else None
                    note = "formula"
                elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
                    val = float(raw)
                elif isinstance(raw, str):
                    s = raw.strip().replace(",", "")
                    if s == "":
                        val = None
                    else:
                        try:
                            val, note = float(s), "number stored as text"
                        except ValueError:
                            val, note = None, f"non-numeric text '{raw.strip()[:30]}'"
                if val is not None and val < 0:
                    note = (note + "; " if note else "") + "negative value"
                font_red = _is_red(_rgb(cf.font.color) if cf.font else "")
                cmt = cf.comment.text.strip() if cf.comment else ""
                out.records.append(CellRecord(
                    path.name, ws.title, r, out.dept_file, section, phase, category,
                    code.full, code.base, code.suffix, code.source, code.note,
                    str(a.value) if a.value is not None else "", text, mth, raw, val, note,
                    font_red, _fill(cf), row_fill, bool(cmt), cmt))
                if note and not note.startswith("formula"):
                    out.issues.append(dict(file=path.name, sheet=ws.title, row=r, severity="Warning",
                                           issue=f"Cell value: {note}", detail=f"{cf.coordinate} '{raw}' ({text[:40]})"))
    return out
